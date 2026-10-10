"""Compare an auxiliary Sber operations export with a broker report.

The input files contain personal financial data and must remain outside Git.
SPDX-License-Identifier: GPL-3.0-or-later
"""

from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from svoibumagi.reports import account_key
from svoibumagi.sber import parse_sber


def _date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.fromisoformat(str(value)).date()


def check(export_path: Path, report_path: Path) -> dict[str, int]:
    workbook = load_workbook(export_path, read_only=True, data_only=True)
    try:
        required = {"Сделки", "Движение ДС"}
        if not required <= set(workbook.sheetnames):
            raise ValueError("Not a Sber operations export")
        trades = list(workbook["Сделки"].iter_rows(values_only=True))
        cash = list(workbook["Движение ДС"].iter_rows(values_only=True))
    finally:
        workbook.close()
    if not trades or tuple(trades[0][:3]) != ("Номер договора", "Номер сделки", "Дата заключения"):
        raise ValueError("Sber trade export columns changed")
    if not cash or tuple(cash[0][:4]) != ("Номер договора", "Дата подачи поручения", "Дата исполнения поручения", "Операция"):
        raise ValueError("Sber cash export columns changed")
    report = parse_sber(report_path)
    if report.account_kind != "brokerage":
        raise ValueError("Auxiliary export must be compared with the matching brokerage account")
    by_id = {trade.broker_id: trade for trade in report.trades}
    if len(by_id) != len(report.trades):
        raise ValueError("Report has duplicate broker trade IDs")
    trade_rows = [row for row in trades[1:] if row[1] not in (None, "")]
    seen: set[str] = set()
    price_scale_ten = 0
    for row in trade_rows:
        contract = str(row[0]).strip()
        if account_key("sber", contract) != report.account_key:
            raise ValueError("Contract in export differs from broker report")
        broker_id = str(row[1]).strip()
        if broker_id in seen:
            raise ValueError("Duplicate trade in operations export")
        seen.add(broker_id)
        trade = by_id.get(broker_id)
        if trade is None:
            raise ValueError("Export trade is missing from broker report")
        direction = {"Покупка": 1, "Продажа": -1}.get(str(row[6]).strip())
        if direction is None or str(row[11]).strip() != "RUB":
            raise ValueError("Unsupported trade direction or currency in export")
        quantity = Decimal(str(row[7])) * direction
        if (trade.traded_at != _date(row[2]) or trade.quantity != quantity
                or trade.amount != Decimal(str(row[10]))):
            raise ValueError("Trade economics differ between export and report")
        if row[4] == "Облигация" and trade.price * 10 == Decimal(str(row[8])):
            price_scale_ten += 1
    withdrawal_rows = [row for row in cash[1:] if row[3] == "Вывод со счёта" and row[10] == "Исполнена"]
    reported_withdrawals = Counter((movement.date, movement.amount) for movement in report.cash_movements
                                   if movement.kind == "withdrawal")
    for row in withdrawal_rows:
        if account_key("sber", str(row[0]).strip()) != report.account_key:
            raise ValueError("Cash export contract differs from broker report")
        key = (_date(row[2]), -Decimal(str(row[5])))
        if reported_withdrawals[key] <= 0:
            raise ValueError("Executed withdrawal is missing from broker report")
        reported_withdrawals[key] -= 1
    return {"matched_trades": len(trade_rows), "matched_withdrawals": len(withdrawal_rows),
            "bond_price_scale_ten": price_scale_ten}


def main() -> None:
    parser = argparse.ArgumentParser(description="Private Sber auxiliary export check")
    parser.add_argument("export", type=Path)
    parser.add_argument("report", type=Path)
    args = parser.parse_args()
    result = check(args.export, args.report)
    print("Matched {matched_trades} trades and {matched_withdrawals} withdrawals; "
          "{bond_price_scale_ten} bond prices use a tenfold scale in the export".format(**result))


if __name__ == "__main__":
    main()
