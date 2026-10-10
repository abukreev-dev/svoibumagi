"""Conservative Alfa XLS/XLSX report parser. SPDX-License-Identifier: GPL-3.0-or-later"""

from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
import zipfile

from .reports import CashFlowSubtotal, CashMovement, CashSnapshot, Location, ParsedReport, PositionSnapshot, SourceTotal, Trade, UnparsedRow, account_key


REQUIRED = {"Динамика позиций", "Завершенные сделки", "Незавершенные сделки", " Движение ДС", "Неторговые операции"}
ADAPTER_VERSION = "alfa-xls-dynamic-positions-v12"
MAX_WORKBOOK_CELLS = 1_000_000
MAX_XLSX_EXPANDED_BYTES = 128_000_000
MAX_XLSX_ENTRIES = 1_000
HEADER_REASON = "Alfa report heading or column labels; no ledger event"
ISIN_CONTINUATION_REASON = "ISIN continuation used by preceding position; no ledger event"
POSITION_FOOTNOTE_REASON = "Alfa position table footnote; no ledger event"
TOTAL_REASON = "Alfa cash movement subtotal; see source_total check"
POSITION_TOTAL_REASON = "Alfa portfolio value subtotal; see source_total checks"
TAX_REASON = "Alfa tax liability summary; arithmetic check only, no ledger event"
TAX_HEADER_REASON = "Alfa tax liability heading; no ledger event"
TRADE_FOOTNOTE_REASON = "Alfa trade table footer note; no ledger event"
EMPTY_TRADE_SHEET_REASON = "Alfa empty trade sheet heading; no ledger event"
CASH_SUMMARY_REASON = "Alfa period cash and portfolio summary; see source_total checks"
CASH_SECTION_REASON = "Alfa cash section label; no ledger event"
NONTRADE_METADATA_REASON = "Alfa market access period metadata; no security operation"
EMPTY_NONTRADE_SHEET_REASON = "Alfa empty non-trading sheet heading; no security operation"
ISIN = re.compile(r"^[A-Z]{2}[A-Z0-9]{9}[0-9]$")
PERIOD = re.compile(r"(\d{2}\.\d{2}\.\d{4})\s*-\s*(\d{2}\.\d{2}\.\d{4})")
TRADE_REF = re.compile(r"(?:Расчеты|Комиссия|НКД) по сделке\s+(.+)$", re.I)
DIVIDEND_TAX = re.compile(r"Удержан налог\s+([\d\s\u00a0]+(?:[.,]\d{1,2})?)\s*руб\.", re.I)


class UnsupportedReport(ValueError):
    pass


def _num(value: object) -> Decimal:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        result = Decimal(str(value))
        if result.is_finite():
            return result
        raise ValueError("not a finite decimal number")
    if isinstance(value, str):
        cleaned = value.strip().replace("\u00a0", "").replace(" ", "").replace(",", ".")
        if cleaned:
            try:
                result = Decimal(cleaned)
            except InvalidOperation:
                pass
            else:
                if result.is_finite():
                    return result
                raise ValueError("not a finite decimal number")
    raise ValueError("not a decimal number")


def _date(value: object, datemode: int = 0) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        import xlrd

        return xlrd.xldate_as_datetime(value, datemode).date()
    if isinstance(value, str):
        m = re.search(r"\b\d{2}\.\d{2}\.\d{4}\b", value)
        if m:
            return datetime.strptime(m.group(), "%d.%m.%Y").date()
    raise ValueError("not a report date")


def _execution_time(value: object) -> str:
    matches = re.findall(r"\b(\d{1,2}):(\d{2}):(\d{2})\b", str(value))
    if len(matches) != 1:
        raise ValueError("trade execution time is absent or ambiguous")
    try:
        return time(*(int(part) for part in matches[0])).isoformat()
    except ValueError as exc:
        raise ValueError("trade execution time is invalid") from exc


def _sheet_rows(path: Path) -> tuple[dict[str, list[list[object]]], int]:
    if path.suffix.lower() == ".xls":
        import xlrd

        try:
            book = xlrd.open_workbook(path, logfile=None, on_demand=True)
        except Exception as exc:
            raise UnsupportedReport("Unreadable XLS") from exc
        try:
            sheets = book.sheets()
            if sum(sheet.nrows * sheet.ncols for sheet in sheets) > MAX_WORKBOOK_CELLS:
                raise UnsupportedReport("XLS exceeds parser cell limit")
            return ({sheet.name: [sheet.row_values(i) for i in range(sheet.nrows)] for sheet in sheets}, book.datemode)
        finally:
            book.release_resources()
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        try:
            with zipfile.ZipFile(path) as archive:
                entries = archive.infolist()
                if (len(entries) > MAX_XLSX_ENTRIES
                        or sum(item.file_size for item in entries) > MAX_XLSX_EXPANDED_BYTES):
                    raise UnsupportedReport("XLSX exceeds parser expansion limit")
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                cells = 0
                sheets = {}
                for sheet in book:
                    if sheet.max_row and sheet.max_column and sheet.max_row * sheet.max_column > MAX_WORKBOOK_CELLS:
                        raise UnsupportedReport("XLSX exceeds parser cell limit")
                    rows = []
                    for row in sheet.iter_rows(values_only=True):
                        cells += len(row)
                        if cells > MAX_WORKBOOK_CELLS:
                            raise UnsupportedReport("XLSX exceeds parser cell limit")
                        rows.append(list(row))
                    sheets[sheet.title] = rows
                return sheets, 0
            finally:
                book.close()
        except UnsupportedReport:
            raise
        except Exception as exc:
            raise UnsupportedReport("Unreadable XLSX") from exc
    raise UnsupportedReport("Only XLS/XLSX files are accepted")


def _get(row: list[object], col: int) -> object:
    return row[col] if col < len(row) and row[col] is not None else ""


def _period(rows: dict[str, list[list[object]]]) -> tuple[date, date]:
    found = set()
    for sheet in rows.values():
        for row in sheet[:10]:
            for value in row:
                match = PERIOD.search(str(value))
                if match:
                    found.add(tuple(datetime.strptime(x, "%d.%m.%Y").date() for x in match.groups()))
    if len(found) != 1:
        raise UnsupportedReport("Report period missing or inconsistent")
    start, end = next(iter(found))
    if start > end:
        raise UnsupportedReport("Report period ends before it starts")
    return start, end


def parse_alfa(path: Path) -> ParsedReport:
    rows, datemode = _sheet_rows(path)
    if not REQUIRED.issubset(rows):
        raise UnsupportedReport("Alfa report sheets not found")
    start, end = _period(rows)
    report = ParsedReport("alfa", ADAPTER_VERSION, start, end)
    agreements = {str(value).strip() for row in rows[" Движение ДС"][:8] for value in row
                  if re.fullmatch(r"№\s*\d+\s+от\s+\d{2}\.\d{2}\.\d{4}", str(value).strip())}
    if len(agreements) != 1:
        raise UnsupportedReport("Alfa agreement identity is missing or ambiguous")
    agreement = next(iter(agreements))
    report.account_key = account_key("alfa", agreement)
    report.account_number = re.search(r"\d+", agreement).group()

    trade_sheet = rows["Завершенные сделки"]
    header = next((i for i, row in enumerate(trade_sheet) if "№ сделки" in str(_get(row, 4)) and "ISIN" in str(_get(row, 12))), None)
    if header is None and any(str(_get(row, 4)).strip().splitlines()[:1] and re.match(r"^\d{5,}$", str(_get(row, 4)).strip().splitlines()[0]) for row in trade_sheet[15:]):
        raise UnsupportedReport("Completed trade columns not recognized")
    hrow = trade_sheet[header] if header is not None else []
    currency_col = next((j for j, val in enumerate(hrow) if "Валюта расчетов" in str(val)), 21)
    commission_col = next((j for j, val in enumerate(hrow) if "Комиссия" in str(val) and "банка" in str(val)), 22)
    for i in range((header + 1) if header is not None else len(trade_sheet), len(trade_sheet)):
        row = trade_sheet[i]
        loc = Location("Завершенные сделки", i + 1)
        trade_id = str(_get(row, 4)).strip().splitlines()
        trade_id = trade_id[0] if trade_id else ""
        trade_like = bool(str(_get(row, 12)).strip()) and bool(str(_get(row, 16)).strip())
        if not trade_id:
            if trade_like:
                report.unparsed.append(UnparsedRow(loc, "trade row without broker ID"))
            continue
        if not re.match(r"^[A-Za-z0-9-]{5,}$", trade_id):
            if trade_like:
                report.unparsed.append(UnparsedRow(loc, "trade row with invalid broker ID"))
            continue
        try:
            qty = _num(_get(row, 16))
            currency = str(_get(row, currency_col)).strip()
            isin = str(_get(row, 12)).strip()
            if currency != "RUB" or not ISIN.fullmatch(isin) or qty == 0:
                raise ValueError("unsupported currency, instrument or quantity")
            traded_at, settled_at = _date(_get(row, 7), datemode), _date(_get(row, 10), datemode)
            if settled_at < traded_at:
                raise ValueError("trade settlement precedes execution")
            if not start <= settled_at <= end:
                raise ValueError("completed trade settlement outside report period")
            report.trades.append(Trade(loc, trade_id, traded_at, settled_at, isin,
                                       str(_get(row, 14)).strip(), qty, _num(_get(row, 18)), _num(_get(row, 19)),
                                       _num(_get(row, commission_col)), currency,
                                       execution_time=_execution_time(_get(row, 7)), execution_time_required=True))
        except ValueError as exc:
            report.unparsed.append(UnparsedRow(loc, f"trade: {exc}"))

    cash_sheet = rows[" Движение ДС"]
    cash_header = next((i for i, row in enumerate(cash_sheet) if "Наименование операции" in str(_get(row, 9))), None)
    if cash_header is None:
        raise UnsupportedReport("Cash movement columns not recognized")
    income_candidates: dict[int, str] = {}
    income_sums = {"dividend": Decimal(0), "coupon": Decimal(0)}
    for i in range(cash_header + 2, len(cash_sheet)):
        row = cash_sheet[i]
        if str(_get(row, 9)).strip() == "Итого:" or str(_get(row, 10)).strip() == "Итого:":
            break
        if str(_get(row, 9)).strip() != "Перевод":
            continue
        comment = str(_get(row, 10)).strip().lower()
        category = ("dividend" if comment.startswith("дивиденды ") else
                    "coupon" if comment.startswith("погашение купона ") else None)
        if category:
            try:
                value = _num(_get(row, 14))
            except ValueError:
                continue
            if value > 0:
                income_candidates[i] = category
                income_sums[category] += value
    income_subtotals: dict[str, list[Decimal]] = {"dividend": [], "coupon": []}
    for i, row in enumerate(cash_sheet, 1):
        label = str(_get(row, 11)).strip()
        category = {"Дивиденды": "dividend", "Купоны": "coupon"}.get(label)
        if category:
            loc = Location(" Движение ДС", i)
            if income_subtotals[category]:
                report.unparsed.append(UnparsedRow(loc, "duplicate income subtotal"))
                continue
            try:
                amount = _num(_get(row, 14))
            except ValueError:
                report.unparsed.append(UnparsedRow(loc, "income subtotal without RUB amount"))
                continue
            income_subtotals[category].append(amount)
            report.cash_flow_subtotals.append(CashFlowSubtotal(loc, category, amount))
    verified_income = {kind for kind, values in income_subtotals.items()
                       if len(values) == 1 and values[0] == income_sums[kind] and income_sums[kind] > 0}
    last_day: date | None = None
    for i in range(cash_header + 2, len(cash_sheet)):
        row = cash_sheet[i]
        loc = Location(" Движение ДС", i + 1)
        if str(_get(row, 9)).strip() == "Итого:" or str(_get(row, 10)).strip() == "Итого:":
            break
        invalid_day = False
        if _get(row, 2) not in ("", None):
            try:
                last_day = _date(_get(row, 2), datemode)
            except ValueError:
                last_day = None
                invalid_day = True
        kind = str(_get(row, 9)).strip()
        if not kind:
            if (last_day and str(_get(row, 14)).strip() not in {"", "0", "0.0"}
                    and (str(_get(row, 6)).strip() or str(_get(row, 26)).strip())):
                report.unparsed.append(UnparsedRow(loc, "cash amount without operation type"))
            continue
        if kind == "Итого:" or kind.startswith("Дивиденды"):
            continue
        if invalid_day:
            report.unparsed.append(UnparsedRow(loc, "cash movement has invalid date"))
            continue
        if not last_day:
            report.unparsed.append(UnparsedRow(loc, "cash movement without date"))
            continue
        if not start <= last_day <= end:
            report.unparsed.append(UnparsedRow(loc, "cash movement date outside report period"))
            continue
        try:
            amount = _num(_get(row, 14))
        except ValueError:
            report.unparsed.append(UnparsedRow(loc, "cash movement without RUB amount"))
            continue
        ref_match = TRADE_REF.match(kind)
        withheld_tax = None
        if kind == "Перевод":
            comment = str(_get(row, 10)).strip().casefold()
            category = income_candidates.get(i) if income_candidates.get(i) in verified_income else None
            if category == "dividend":
                matches = DIVIDEND_TAX.findall(comment)
                if len(matches) == 1:
                    try:
                        withheld_tax = _num(matches[0])
                        if not withheld_tax.is_finite() or withheld_tax < 0:
                            raise ValueError("invalid dividend withholding")
                    except ValueError:
                        report.unparsed.append(UnparsedRow(loc, "dividend withholding amount is invalid"))
                        withheld_tax = None
                else:
                    report.unparsed.append(UnparsedRow(loc, "dividend withholding amount is absent or ambiguous"))
            if category is None and (comment.startswith('из ао "альфа-банк"') or comment.startswith("с субсчета ")) and amount > 0:
                category = "deposit"
                if comment.startswith("с субсчета "):
                    report.unparsed.append(UnparsedRow(loc, "cash deposit from subaccount; portfolio counterparty not verified"))
            elif category is None and comment.startswith("возврат средств по дог.") and amount < 0:
                category = "withdrawal"
            if category is None:
                category = "transfer_unclassified"
                reason = ("income transfer lacks matching report subtotal" if i in income_candidates
                          else "cash transfer source/counterparty not verified")
                report.unparsed.append(UnparsedRow(loc, reason))
        elif kind == "НДФЛ":
            category = "tax_withheld"
        elif kind == "Комиссия":
            category = "fee"
        elif kind == "НКД по сделке":
            category = "accrued_interest"
        elif ref_match:
            category = "trade_cash" if kind.startswith("Расчеты") else ("trade_fee" if kind.startswith("Комиссия") else "accrued_interest")
        else:
            report.unparsed.append(UnparsedRow(loc, "unknown cash movement type"))
            continue
        if category == "dividend" and withheld_tax is not None:
            report.cash_movements.append(CashMovement(loc, last_day, "dividend", amount + withheld_tax,
                                                       "RUB", withheld_tax=withheld_tax))
            if withheld_tax:
                report.cash_movements.append(CashMovement(loc, last_day, "tax_withheld", -withheld_tax,
                                                           "RUB", component="dividend_tax"))
        else:
            report.cash_movements.append(CashMovement(loc, last_day, category, amount, "RUB",
                                                       ref_match.group(1).strip() if ref_match else None))

    cash_total_rows = [(i, row) for i, row in enumerate(cash_sheet, 1)
                       if str(_get(row, 10)).strip() == "Итого:" and i > cash_header + 1]
    if len(cash_total_rows) == 1:
        row_number, row = cash_total_rows[0]
        loc = Location(" Движение ДС", row_number)
        try:
            reported_total = _num(_get(row, 14))
            calculated_total = sum((movement.amount for movement in report.cash_movements), Decimal(0))
            report.source_totals.append(SourceTotal(loc, f" Движение ДС:{row_number}:field:14",
                                                    calculated_total, reported_total))
            report.excluded.append(UnparsedRow(loc, TOTAL_REASON))
        except ValueError:
            report.unparsed.append(UnparsedRow(loc, "cash movement subtotal without RUB amount"))
    elif len(cash_total_rows) > 1:
        report.unparsed.extend(UnparsedRow(Location(" Движение ДС", row_number),
                                           "ambiguous cash movement subtotal") for row_number, _ in cash_total_rows)

    subtotal_labels = {"Внесено по распоряжению Клиента": "deposit",
                       "Списано по распоряжению Клиента": "withdrawal",
                       "Списано по тарифам Банка": "bank_fee"}
    seen_subtotals: set[str] = set()
    for i, row in enumerate(cash_sheet, 1):
        label = str(_get(row, 11)).strip()
        category = subtotal_labels.get(label)
        if category is None:
            continue
        loc = Location(" Движение ДС", i)
        if category in seen_subtotals:
            report.unparsed.append(UnparsedRow(loc, "duplicate cash flow subtotal"))
            continue
        seen_subtotals.add(category)
        try:
            amount = _num(_get(row, 14))
        except ValueError:
            report.unparsed.append(UnparsedRow(loc, "cash flow subtotal without RUB amount"))
            continue
        report.cash_flow_subtotals.append(CashFlowSubtotal(loc, category, amount))

    positions = rows["Динамика позиций"]
    position_header = next((i for i, row in enumerate(positions) if str(_get(row, 2)).strip() == "Инструмент" and str(_get(row, 6)).strip().startswith("Актив")), None)
    if position_header is None:
        raise UnsupportedReport("Position columns not recognized")
    active_category: str | None = None
    security_categories = {"Акции", "Облигации", "Прочее", "Паи"}
    for i in range(position_header + 1, len(positions)):
        row = positions[i]
        category = str(_get(row, 2)).strip()
        name = str(_get(row, 6)).strip()
        loc = Location("Динамика позиций", i + 1)
        if category == "Валюта" and name == "RUB":
            active_category = None
            try:
                report.cash_snapshots.append(CashSnapshot(loc, end, _num(_get(row, 12)), _num(_get(row, 15)), "RUB"))
            except ValueError:
                report.unparsed.append(UnparsedRow(loc, "invalid RUB position"))
            continue
        if category == "Валюта" and name != "RUB":
            active_category = None
            try:
                if _num(_get(row, 12)) != 0 or _num(_get(row, 15)) != 0:
                    report.unparsed.append(UnparsedRow(loc, "non-RUB cash balance"))
            except ValueError:
                report.unparsed.append(UnparsedRow(loc, "unrecognized currency balance"))
            continue
        if category in security_categories:
            active_category = category
        elif category:
            active_category = None
        if not name or not active_category:
            continue
        next_name = str(_get(positions[i + 1], 6)).strip() if i + 1 < len(positions) else ""
        if not ISIN.fullmatch(next_name) and category == "" and not any(str(_get(row, col)).strip() for col in (12, 15)):
            continue
        try:
            isin = next_name if ISIN.fullmatch(next_name) else None
            if not isin:
                report.unparsed.append(UnparsedRow(loc, "position instrument has no ISIN"))
            report.positions.append(PositionSnapshot(loc, end, isin, name, _num(_get(row, 12)), _num(_get(row, 15))))
        except ValueError:
            report.unparsed.append(UnparsedRow(loc, "invalid position quantity"))

    position_totals = [(i, row) for i, row in enumerate(positions, 1)
                       if str(_get(row, 6)).strip() == "Стоимость всех позиций, руб."]
    if len(position_totals) == 1:
        row_number, row = position_totals[0]
        loc = Location("Динамика позиций", row_number)
        detail_rows = [positions[item.location.row - 1]
                       for item in report.positions + report.cash_snapshots
                       if item.location.sheet == "Динамика позиций"]
        try:
            for field in (14, 17):
                calculated = sum((_num(_get(detail, field)) if _get(detail, field) != "" else Decimal(0)
                                  for detail in detail_rows), Decimal(0))
                report.source_totals.append(SourceTotal(loc, f"Динамика позиций:{row_number}:field:{field}",
                                                        calculated, _num(_get(row, field))))
            report.excluded.append(UnparsedRow(loc, POSITION_TOTAL_REASON))
        except ValueError:
            report.unparsed.append(UnparsedRow(loc, "portfolio value subtotal has unreadable amount"))
    elif len(position_totals) > 1:
        report.unparsed.extend(UnparsedRow(Location("Динамика позиций", row_number),
                                           "ambiguous portfolio value subtotal") for row_number, _ in position_totals)

    tax_labels = ("НДФЛ исчислено, руб.", "Удержано, руб.", "К доплате, руб.")
    tax_rows = {label: [(i, row) for i, row in enumerate(positions, 1)
                        if str(_get(row, 14)).strip() == label] for label in tax_labels}
    if all(len(tax_rows[label]) == 1 for label in tax_labels):
        try:
            amounts = {label: _num(_get(tax_rows[label][0][1], 17)) for label in tax_labels}
            due_row = tax_rows["К доплате, руб."][0][0]
            report.source_totals.append(SourceTotal(Location("Динамика позиций", due_row),
                                                    f"Динамика позиций:{due_row}:tax_due",
                                                    amounts[tax_labels[0]] - amounts[tax_labels[1]],
                                                    amounts[tax_labels[2]]))
            for label in tax_labels:
                report.excluded.append(UnparsedRow(Location("Динамика позиций", tax_rows[label][0][0]), TAX_REASON))
        except ValueError:
            report.unparsed.extend(UnparsedRow(Location("Динамика позиций", row_number),
                                               "tax summary has unreadable amount")
                                   for entries in tax_rows.values() for row_number, _ in entries)
    elif any(tax_rows.values()):
        report.unparsed.extend(UnparsedRow(Location("Динамика позиций", row_number),
                                           "tax summary is incomplete or ambiguous")
                               for entries in tax_rows.values() for row_number, _ in entries)
    for row_number, row in enumerate(positions, 1):
        if str(_get(row, 14)).strip() in {"Предварительный расчет обязательств по НДФЛ",
                                         "Расчет обязательств по НДФЛ"}:
            report.excluded.append(UnparsedRow(Location("Динамика позиций", row_number), TAX_HEADER_REASON))

    rub_cash = [item for item in report.cash_snapshots if item.location.sheet == "Динамика позиций"
                and item.currency == "RUB"]
    if len(rub_cash) == 1 and len(position_totals) == 1:
        for label, cash_amount, total_field in (("На начало периода", rub_cash[0].start_amount, 14),
                                                 ("На конец периода", rub_cash[0].end_amount, 17)):
            summary_headers = [i for i, row in enumerate(cash_sheet)
                               if str(_get(row, 11)).strip() == label]
            if len(summary_headers) != 1:
                continue
            header_index = summary_headers[0]
            row_index = header_index + 2
            loc = Location(" Движение ДС", row_index + 1)
            try:
                if str(_get(cash_sheet[header_index + 1], 14)).strip() != "RUB":
                    raise ValueError("cash summary currency is not RUB")
                source = cash_sheet[row_index]
                for field, calculated in ((14, cash_amount), (16, cash_amount),
                                          (19, _num(_get(position_totals[0][1], total_field)))):
                    report.source_totals.append(SourceTotal(loc, f" Движение ДС:{loc.row}:field:{field}",
                                                            calculated, _num(_get(source, field))))
                report.excluded.append(UnparsedRow(loc, CASH_SUMMARY_REASON))
            except (IndexError, ValueError):
                report.unparsed.append(UnparsedRow(loc, "cash period summary is unreadable"))
    for row_number, row in enumerate(cash_sheet, 1):
        label = str(_get(row, 11)).strip()
        if (label == "На конец периода" or label == "В том числе" or
                str(_get(row, 14)).strip() == "RUB" and
                (str(_get(row, 17)).strip() == "ден. позиций" or
                 str(_get(row, 16)).strip() == "денежных")):
            if row_number > cash_header + 1:
                report.excluded.append(UnparsedRow(Location(" Движение ДС", row_number), CASH_SECTION_REASON))

    # Other sections may contain transfers/corporate actions. Until classified,
    # they remain visible and prevent a misleading verified status.
    nontrade = rows["Неторговые операции"]
    nh = next((i for i, row in enumerate(nontrade)
               if "Наименование операции" in str(_get(row, 10)) or
               "Наименование операции" in str(_get(row, 9))), None)
    if nh is not None:
        for i in range(nh + 1, len(nontrade)):
            row = nontrade[i]
            if (_get(row, 10) or _get(row, 9)) and _get(row, 2):
                report.unparsed.append(UnparsedRow(Location("Неторговые операции", i + 1), "non-trading security operation requires classification"))
    pending = rows["Незавершенные сделки"]
    ph = next((i for i, row in enumerate(pending) if "№ сделки" in str(_get(row, 4))), None)
    if ph is not None:
        for i in range(ph + 1, len(pending)):
            ident = str(_get(pending[i], 4)).strip().splitlines()
            if ident and re.match(r"^[A-Za-z0-9-]{5,}$", ident[0]):
                report.excluded.append(UnparsedRow(Location("Незавершенные сделки", i + 1), "unsettled trade"))

    classified = {(item.location.sheet, item.location.row)
                  for group in (report.trades, report.cash_movements, report.positions, report.cash_snapshots,
                                report.cash_flow_subtotals, report.unparsed, report.excluded)
                  for item in group}
    for sheet_name, header_index in (("Завершенные сделки", header), (" Движение ДС", cash_header),
                                     ("Динамика позиций", position_header), ("Неторговые операции", nh),
                                     ("Незавершенные сделки", ph)):
        if header_index is None:
            continue
        for index, row in enumerate(rows[sheet_name][:header_index + 1]):
            loc = Location(sheet_name, index + 1)
            if (loc.sheet, loc.row) in classified or not any(str(value).strip() for value in row if value is not None):
                continue
            report.excluded.append(UnparsedRow(loc, HEADER_REASON))
    position_locations = {item.location.row for item in report.positions}
    footnote_prefixes = ("¹ Активы:", "² Стоимость позиций:", "³ Плановое количество",
                         "⁴ Места учета", "⁵ грамм:", "законодательством Российской Федерации.",
                         "Срочн. рынок —", "Euroclear —")
    for index, row in enumerate(positions[position_header + 1:], position_header + 1):
        loc = Location("Динамика позиций", index + 1)
        if (loc.sheet, loc.row) in classified:
            continue
        filled = [(column, str(value).strip()) for column, value in enumerate(row)
                  if value is not None and str(value).strip()]
        if len(filled) != 1 or filled[0][0] != 6:
            continue
        value = filled[0][1]
        if ISIN.fullmatch(value) and index in position_locations:
            report.excluded.append(UnparsedRow(loc, ISIN_CONTINUATION_REASON))
        elif value.startswith(footnote_prefixes):
            report.excluded.append(UnparsedRow(loc, POSITION_FOOTNOTE_REASON))
    for sheet_name, header_index, last_data in (
        ("Завершенные сделки", header,
         max((item.location.row for item in report.trades if item.location.sheet == "Завершенные сделки"), default=0)),
        ("Незавершенные сделки", ph,
         max((item.location.row for item in report.excluded
              if item.location.sheet == "Незавершенные сделки" and item.reason == "unsettled trade"), default=0)),
    ):
        if header_index is None:
            continue
        for index, row in enumerate(rows[sheet_name]):
            if index <= header_index or index + 1 <= last_data:
                continue
            filled = [(column, str(value).strip()) for column, value in enumerate(row)
                      if value is not None and str(value).strip()]
            if len(filled) == 1 and filled[0][0] == 11:
                loc = Location(sheet_name, index + 1)
                if (loc.sheet, loc.row) not in classified:
                    report.excluded.append(UnparsedRow(loc, TRADE_FOOTNOTE_REASON))
    for sheet_name, header_index in (("Завершенные сделки", header), ("Незавершенные сделки", ph)):
        if header_index is not None:
            continue
        for index, row in enumerate(rows[sheet_name]):
            trade_id = str(_get(row, 4)).strip().splitlines()[:1]
            if ((trade_id and re.fullmatch(r"[A-Za-z0-9-]{5,}", trade_id[0])) or
                    ISIN.fullmatch(str(_get(row, 12)).strip()) or
                    re.fullmatch(r"[+-]?\d+(?:[.,]\d+)?", str(_get(row, 16)).strip())):
                raise UnsupportedReport(f"{sheet_name} has trade-like data without recognized columns")
            if any(str(value).strip() for value in row if value is not None):
                report.excluded.append(UnparsedRow(Location(sheet_name, index + 1), EMPTY_TRADE_SHEET_REASON))
    if nh is not None:
        for index, row in enumerate(nontrade[nh + 1:], nh + 1):
            loc = Location("Неторговые операции", index + 1)
            if any(item.location == loc for item in report.unparsed):
                continue
            first = str(_get(row, 0)).strip()
            marker = str(_get(row, 10)).strip()
            period_date = False
            if marker == "-" and first and all(_get(row, column) == "" for column in range(len(row))
                                                     if column not in {0, 10}):
                try:
                    _date(_get(row, 0), datemode)
                    period_date = True
                except ValueError:
                    pass
            if (first in {"Периоды использования единого финансового рынка",
                          "Единый финансовый рынок (на конец периода)"} or
                    first == "Начало периода ¹" and marker == "Окончание периода ¹" or
                    first.startswith("¹время указано по часовому поясу") or
                    period_date):
                report.excluded.append(UnparsedRow(loc, NONTRADE_METADATA_REASON))
    else:
        for row_number, row in enumerate(nontrade, 1):
            if _get(row, 2) and (_get(row, 9) or _get(row, 10)):
                raise UnsupportedReport("Alfa non-trading sheet has operation-like data without recognized columns")
            if any(value is not None and str(value).strip() for value in row):
                report.excluded.append(UnparsedRow(Location("Неторговые операции", row_number),
                                                   EMPTY_NONTRADE_SHEET_REASON))

    classified = {(item.location.sheet, item.location.row)
                  for group in (report.trades, report.cash_movements, report.positions, report.cash_snapshots,
                                report.cash_flow_subtotals, report.unparsed, report.excluded)
                  for item in group}
    for sheet_name in REQUIRED:
        for row_number, row in enumerate(rows[sheet_name], 1):
            if (sheet_name, row_number) not in classified and any(
                    value is not None and str(value).strip() for value in row):
                report.unparsed.append(UnparsedRow(Location(sheet_name, row_number),
                                                   "unrecognized nonempty Alfa source row"))
    return report
