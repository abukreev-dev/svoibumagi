"""Independently inventory all nonempty rows in recognized private Alfa sheets.

SPDX-License-Identifier: GPL-3.0-or-later
"""

from __future__ import annotations

from collections import Counter
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import re
import sys

from svoibumagi.alfa import CASH_SECTION_REASON, CASH_SUMMARY_REASON, EMPTY_TRADE_SHEET_REASON, HEADER_REASON, ISIN_CONTINUATION_REASON, NONTRADE_METADATA_REASON, POSITION_FOOTNOTE_REASON, POSITION_TOTAL_REASON, TAX_HEADER_REASON, TAX_REASON, TRADE_FOOTNOTE_REASON, TOTAL_REASON, parse_alfa


ISIN = re.compile(r"[A-Z]{2}[A-Z0-9]{9}[0-9]")
REFERENCE = re.compile(r"[A-Za-z0-9-]{5,}")
SHEETS = {"Динамика позиций", "Завершенные сделки", "Незавершенные сделки",
          " Движение ДС", "Неторговые операции"}


def _read(path: Path) -> dict[str, list[list[object]]]:
    if path.suffix.lower() == ".xls":
        import xlrd

        book = xlrd.open_workbook(path, logfile=None)
        return {sheet.name: [sheet.row_values(i) for i in range(sheet.nrows)] for sheet in book.sheets()}
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    return {sheet.title: [list(row) for row in sheet.iter_rows(values_only=True)] for sheet in book}


def _value(row: list[object], column: int) -> str:
    return str(row[column]).strip() if column < len(row) and row[column] is not None else ""


def _economic(sheet: str, rows: list[list[object]], index: int) -> bool:
    row = rows[index]
    value = lambda col: _value(row, col)
    if sheet == "Завершенные сделки":
        return bool(REFERENCE.fullmatch(value(4).splitlines()[0] if value(4) else "")
                    and ISIN.fullmatch(value(12)))
    if sheet == "Незавершенные сделки":
        return bool(REFERENCE.fullmatch(value(4).splitlines()[0] if value(4) else "")
                    and value(4) != "№ сделки")
    if sheet == "Динамика позиций":
        if value(2) == "Валюта" and value(6) == "RUB":
            return True
        return bool(value(6) and index + 1 < len(rows) and ISIN.fullmatch(_value(rows[index + 1], 6)))
    if sheet == " Движение ДС":
        movement = bool(value(9) and value(14) and value(9) not in {"Итого:", "Наименование операции"}
                        and not value(9).startswith("Дивиденды"))
        subtotal = value(11) in {"Внесено по распоряжению Клиента", "Списано по распоряжению Клиента",
                                 "Дивиденды", "Купоны"}
        return movement or subtotal
    return bool(value(2) and (value(9) or value(10)) and
                "Наименование операции" not in {value(9), value(10)})


def check(registry_path: Path) -> tuple[int, Counter[str]]:
    registry = json.loads(registry_path.read_text())
    checked = 0
    candidates: Counter[str] = Counter()
    all_rows: Counter[str] = Counter()
    missing = []
    for item in registry["reports"]:
        if item["broker"] != "alfa":
            continue
        path = Path(item["path"])
        with path.open("rb") as stream:
            if hashlib.file_digest(stream, "sha256").hexdigest() != item["sha256"]:
                raise ValueError("Registered Alfa report changed")
        report = parse_alfa(path)
        classified = {(row.location.sheet, row.location.row)
                      for group in (report.positions, report.cash_snapshots, report.cash_movements,
                                    report.cash_flow_subtotals, report.trades, report.unparsed, report.excluded)
                      for row in group}
        headers = {(row.location.sheet, row.location.row) for row in report.excluded
                   if row.reason == HEADER_REASON}
        references = {(row.location.sheet, row.location.row) for row in report.excluded
                      if row.reason == ISIN_CONTINUATION_REASON}
        footnotes = {(row.location.sheet, row.location.row) for row in report.excluded
                     if row.reason == POSITION_FOOTNOTE_REASON}
        position_rows = {row.location.row for row in report.positions}
        cash_totals = {(row.location.sheet, row.location.row): row for row in report.source_totals
                       if row.location.sheet == " Движение ДС"}
        total_rows = {(row.location.sheet, row.location.row) for row in report.excluded
                      if row.reason == TOTAL_REASON}
        portfolio_totals = {(row.location.sheet, row.location.row, int(row.key.rsplit(":", 1)[1])): row
                            for row in report.source_totals if row.location.sheet == "Динамика позиций"
                            and ":field:" in row.key}
        portfolio_rows = {(row.location.sheet, row.location.row) for row in report.excluded
                          if row.reason == POSITION_TOTAL_REASON}
        position_details = {row.location.row for row in report.positions + report.cash_snapshots
                            if row.location.sheet == "Динамика позиций"}
        tax_classified = {(row.location.sheet, row.location.row) for row in report.excluded
                          if row.reason == TAX_REASON}
        tax_headings = {(row.location.sheet, row.location.row) for row in report.excluded
                        if row.reason == TAX_HEADER_REASON}
        trade_footnotes = {(row.location.sheet, row.location.row) for row in report.excluded
                           if row.reason == TRADE_FOOTNOTE_REASON}
        empty_trade_rows = {(row.location.sheet, row.location.row) for row in report.excluded
                            if row.reason == EMPTY_TRADE_SHEET_REASON}
        nontrade_metadata = {(row.location.sheet, row.location.row) for row in report.excluded
                             if row.reason == NONTRADE_METADATA_REASON}
        cash_summary_rows = {(row.location.sheet, row.location.row) for row in report.excluded
                             if row.reason == CASH_SUMMARY_REASON}
        cash_section_rows = {(row.location.sheet, row.location.row) for row in report.excluded
                             if row.reason == CASH_SECTION_REASON}
        cash_summary_values = {(row.location.sheet, row.location.row, int(row.key.rsplit(":", 1)[1])): row
                               for row in report.source_totals if row.location.sheet == " Движение ДС"
                               and ":field:" in row.key}
        sheets = _read(path)
        for sheet, rows in sheets.items():
            if sheet not in SHEETS:
                continue
            if sheet == "Динамика позиций":
                header = next((i for i, row in enumerate(rows) if _value(row, 2) == "Инструмент" and
                               _value(row, 6).startswith("Актив")), None)
            elif sheet == "Завершенные сделки":
                header = next((i for i, row in enumerate(rows) if "№ сделки" in _value(row, 4) and
                               "ISIN" in _value(row, 12)), None)
            elif sheet == "Незавершенные сделки":
                header = next((i for i, row in enumerate(rows) if "№ сделки" in _value(row, 4)), None)
            elif sheet == " Движение ДС":
                header = next((i for i, row in enumerate(rows) if "Наименование операции" in _value(row, 9)), None)
            else:
                header = next((i for i, row in enumerate(rows)
                               if "Наименование операции" in _value(row, 10) or
                               "Наименование операции" in _value(row, 9)), None)
            if sheet == "Завершенные сделки":
                last_trade = max((row.location.row for row in report.trades
                                  if row.location.sheet == sheet), default=0)
            elif sheet == "Незавершенные сделки":
                last_trade = max((row.location.row for row in report.excluded
                                  if row.location.sheet == sheet and row.reason == "unsettled trade"), default=0)
            else:
                last_trade = 0
            for index in range(len(rows)):
                if sheet in {"Завершенные сделки", "Незавершенные сделки"} and header is None and any(
                        _value(rows[index], col) for col in range(len(rows[index]))):
                    candidates[sheet] += 1
                    if (sheet, index + 1) not in empty_trade_rows:
                        missing.append((item["sha256"][:12], sheet, index + 1, "empty trade sheet"))
                if _economic(sheet, rows, index):
                    candidates[sheet] += 1
                    if (sheet, index + 1) not in classified:
                        missing.append((item["sha256"][:12], sheet, index + 1))
                if header is not None and index <= header and any(_value(rows[index], col) for col in range(len(rows[index]))):
                    candidates[sheet] += 1
                    if (sheet, index + 1) not in headers and (sheet, index + 1) not in cash_summary_rows:
                        missing.append((item["sha256"][:12], sheet, index + 1, "header"))
                if sheet == " Движение ДС" and header is not None and index > header:
                    row = rows[index]
                    section_label = (_value(row, 11) in {"На конец периода", "В том числе"} or
                                     _value(row, 14) == "RUB" and
                                     (_value(row, 17) == "ден. позиций" or _value(row, 16) == "денежных"))
                    if section_label:
                        candidates[sheet] += 1
                        if (sheet, index + 1) not in cash_section_rows:
                            missing.append((item["sha256"][:12], sheet, index + 1, "cash section label"))
                if sheet in {"Завершенные сделки", "Незавершенные сделки"} and header is not None and index > header and index + 1 > last_trade:
                    filled = [col for col in range(len(rows[index])) if _value(rows[index], col)]
                    if filled == [11]:
                        candidates[sheet] += 1
                        if (sheet, index + 1) not in trade_footnotes:
                            missing.append((item["sha256"][:12], sheet, index + 1, "trade footer note"))
                if sheet == "Динамика позиций" and header is not None and index > header:
                    filled = [(col, _value(rows[index], col)) for col in range(len(rows[index]))
                              if _value(rows[index], col)]
                    if len(filled) == 1 and filled[0][0] == 6:
                        value = filled[0][1]
                        location = (sheet, index + 1)
                        if ISIN.fullmatch(value) and index in position_rows:
                            candidates[sheet] += 1
                            if location not in references:
                                missing.append((item["sha256"][:12], sheet, index + 1, "ISIN continuation"))
                        elif value.startswith(("¹ Активы:", "² Стоимость позиций:", "³ Плановое количество",
                                               "⁴ Места учета", "⁵ грамм:",
                                               "законодательством Российской Федерации.", "Срочн. рынок —",
                                               "Euroclear —")):
                            candidates[sheet] += 1
                            if location not in footnotes:
                                missing.append((item["sha256"][:12], sheet, index + 1, "footnote"))
                if sheet == " Движение ДС" and header is not None and index > header and _value(rows[index], 10) == "Итого:":
                    candidates[sheet] += 1
                    location = (sheet, index + 1)
                    total = cash_totals.get(location)
                    try:
                        detail = [Decimal(_value(row, 14).replace(" ", "").replace(",", "."))
                                  for row in rows[header + 1:index]
                                  if _value(row, 9) and _value(row, 14)]
                        calculated = sum(detail, Decimal(0))
                        reported = Decimal(_value(rows[index], 14).replace(" ", "").replace(",", "."))
                    except (InvalidOperation, ValueError):
                        missing.append((item["sha256"][:12], sheet, index + 1, "unreadable total"))
                        continue
                    if (location not in total_rows or total is None or total.calculated != calculated or
                            total.reported != reported or abs(calculated - reported) > Decimal("0.01")):
                        missing.append((item["sha256"][:12], sheet, index + 1, "cash total mismatch"))
                if sheet == "Динамика позиций" and _value(rows[index], 6) == "Стоимость всех позиций, руб.":
                    candidates[sheet] += 1
                    location = (sheet, index + 1)
                    for field in (14, 17):
                        try:
                            calculated = sum((Decimal(_value(rows[row_number - 1], field).replace(" ", "").replace(",", "."))
                                              if _value(rows[row_number - 1], field) else Decimal(0)
                                              for row_number in position_details), Decimal(0))
                            reported = Decimal(_value(rows[index], field).replace(" ", "").replace(",", "."))
                        except (InvalidOperation, ValueError):
                            missing.append((item["sha256"][:12], sheet, index + 1, field, "unreadable"))
                            continue
                        total = portfolio_totals.get((*location, field))
                        if (location not in portfolio_rows or total is None or total.calculated != calculated or
                                total.reported != reported or abs(calculated - reported) > Decimal("0.01")):
                            missing.append((item["sha256"][:12], sheet, index + 1, field, "portfolio total mismatch"))
                if sheet == "Динамика позиций" and _value(rows[index], 14) in {
                    "Предварительный расчет обязательств по НДФЛ", "Расчет обязательств по НДФЛ"}:
                    candidates[sheet] += 1
                    if (sheet, index + 1) not in tax_headings:
                        missing.append((item["sha256"][:12], sheet, index + 1, "tax heading"))
                if sheet == "Динамика позиций" and _value(rows[index], 14) in {
                    "НДФЛ исчислено, руб.", "Удержано, руб.", "К доплате, руб."}:
                    candidates[sheet] += 1
                    if (sheet, index + 1) not in tax_classified:
                        missing.append((item["sha256"][:12], sheet, index + 1, "tax summary"))
                if sheet == "Неторговые операции" and header is not None and index > header:
                    first = _value(rows[index], 0)
                    marker = _value(rows[index], 10)
                    known = (first in {"Периоды использования единого финансового рынка",
                                       "Единый финансовый рынок (на конец периода)"} or
                             first == "Начало периода ¹" and marker == "Окончание периода ¹" or
                             first.startswith("¹время указано по часовому поясу") or
                             marker == "-" and first.replace(".", "", 1).isdigit() and
                             all(not _value(rows[index], col) for col in range(len(rows[index]))
                                 if col not in {0, 10}))
                    if known and (sheet, index + 1) not in nontrade_metadata:
                        missing.append((item["sha256"][:12], sheet, index + 1, "market access metadata"))
                if any(_value(rows[index], col) for col in range(len(rows[index]))):
                    all_rows[sheet] += 1
                    if (sheet, index + 1) not in classified:
                        missing.append((item["sha256"][:12], sheet, index + 1, "unclassified nonempty row"))
            if sheet == "Динамика позиций":
                tax = {_value(row, 14): (i, row) for i, row in enumerate(rows, 1)
                       if _value(row, 14) in {"НДФЛ исчислено, руб.", "Удержано, руб.", "К доплате, руб."}}
                if len(tax) == 3:
                    try:
                        assessed = Decimal(_value(tax["НДФЛ исчислено, руб."][1], 17).replace(" ", "").replace(",", "."))
                        withheld = Decimal(_value(tax["Удержано, руб."][1], 17).replace(" ", "").replace(",", "."))
                        due = Decimal(_value(tax["К доплате, руб."][1], 17).replace(" ", "").replace(",", "."))
                    except (InvalidOperation, ValueError):
                        missing.append((item["sha256"][:12], sheet, "tax arithmetic unreadable"))
                    else:
                        due_row = tax["К доплате, руб."][0]
                        source = next((row for row in report.source_totals
                                       if row.location.sheet == sheet and row.location.row == due_row and
                                       row.key.endswith("tax_due")), None)
                        if source is None or source.calculated != assessed - withheld or source.reported != due or assessed - withheld != due:
                            missing.append((item["sha256"][:12], sheet, due_row, "tax arithmetic mismatch"))
            if sheet == " Движение ДС":
                position_rows = sheets["Динамика позиций"]
                rub = [row for row in position_rows if _value(row, 2) == "Валюта" and _value(row, 6) == "RUB"]
                portfolio = [row for row in position_rows if _value(row, 6) == "Стоимость всех позиций, руб."]
                if len(rub) != 1 or len(portfolio) != 1:
                    missing.append((item["sha256"][:12], sheet, "cash summary source rows"))
                else:
                    for label, cash_field, value_field in (("На начало периода", 12, 14),
                                                           ("На конец периода", 15, 17)):
                        heads = [i for i, row in enumerate(rows) if _value(row, 11) == label]
                        if len(heads) != 1 or heads[0] + 2 >= len(rows):
                            missing.append((item["sha256"][:12], sheet, label, "summary heading"))
                            continue
                        row_index = heads[0] + 2
                        location = (sheet, row_index + 1)
                        if label == "На конец периода":
                            candidates[sheet] += 1
                        try:
                            cash_value = Decimal(_value(rub[0], cash_field).replace(" ", "").replace(",", "."))
                            portfolio_value = Decimal(_value(portfolio[0], value_field).replace(" ", "").replace(",", "."))
                            expected = {14: cash_value, 16: cash_value, 19: portfolio_value}
                        except (InvalidOperation, ValueError):
                            missing.append((item["sha256"][:12], sheet, label, "unreadable summary source"))
                            continue
                        for field, calculated in expected.items():
                            source = cash_summary_values.get((*location, field))
                            try:
                                reported = Decimal(_value(rows[row_index], field).replace(" ", "").replace(",", "."))
                            except InvalidOperation:
                                reported = None
                            if (location not in cash_summary_rows or source is None or reported is None or
                                    source.calculated != calculated or source.reported != reported or
                                    abs(calculated - reported) > Decimal("0.01")):
                                missing.append((item["sha256"][:12], sheet, row_index + 1, field, "cash summary mismatch"))
        checked += 1
    if missing:
        raise ValueError(f"Unclassified or mismatched Alfa rows: {len(missing)}; first locations: {missing[:8]}")
    return checked, all_rows


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("usage: check_alfa_source_rows.py PRIVATE_REGISTRY")
    count, totals = check(Path(sys.argv[1]))
    print(f"Checked {count} Alfa reports and {sum(totals.values())} significant source rows; "
          f"unclassified: 0")
