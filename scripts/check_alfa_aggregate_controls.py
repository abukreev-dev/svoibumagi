"""Check private Alfa aggregate controls against XLS/XLSX cells and parser output.

This is an aggregate cross-check, not a replacement for row-by-row manual controls.
SPDX-License-Identifier: GPL-3.0-or-later
"""

from collections import defaultdict
from decimal import Decimal
import hashlib
from io import StringIO
import json
from pathlib import Path
import re
import sys

from svoibumagi.alfa import parse_alfa


ISIN = re.compile(r"[A-Z]{2}[A-Z0-9]{9}[0-9]")


def _number(value: object) -> Decimal:
    return Decimal(str(value))


def _source_totals(path: Path) -> dict:
    if path.suffix.lower() == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path), logfile=StringIO())
        sheets = {sheet.name: [sheet.row_values(index) for index in range(sheet.nrows)]
                  for sheet in book.sheets()}
    elif path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        sheets = {sheet.title: [list(row) for row in sheet.iter_rows(values_only=True)]
                  for sheet in book}
        book.close()
    else:
        raise ValueError("Aggregate control supports XLS and XLSX sources")
    trades_sheet = sheets["Завершенные сделки"]
    cash_sheet = sheets[" Движение ДС"]
    positions_sheet = sheets["Динамика позиций"]
    trades = [row for row in trades_sheet if ISIN.fullmatch(str(row[12]))]
    movements = [row for row in cash_sheet
                 if str(row[9]).startswith(("Расчеты по сделке", "Комиссия по сделке"))
                 or str(row[9]).strip() == "Комиссия"]
    quantities = defaultdict(Decimal)
    for row in trades:
        quantities[row[12]] += _number(row[16])
    positions = {}
    for index in range(1, len(positions_sheet)):
        isin = str(positions_sheet[index][6])
        if ISIN.fullmatch(isin):
            previous = positions_sheet[index - 1]
            positions[isin] = [_number(previous[12]), _number(previous[15])]

    def cash_balance(label: str) -> Decimal:
        matches = [index for index in range(len(cash_sheet) - 2)
                   if str(cash_sheet[index][11]).strip() == label]
        if len(matches) != 1:
            raise ValueError(f"Expected one cash balance heading: {label}")
        return _number(cash_sheet[matches[0] + 2][14])

    def cash_subtotal(label: str) -> Decimal:
        matches = [row for row in cash_sheet if str(row[11]).strip() == label]
        if len(matches) != 1:
            raise ValueError(f"Expected one cash subtotal: {label}")
        return _number(matches[0][14])

    return {
        "trade_count": len(trades), "cash_row_count": len(movements),
        "buy_amount": sum((_number(row[19]) for row in trades if _number(row[16]) > 0), Decimal(0)),
        "sell_amount": sum((_number(row[19]) for row in trades if _number(row[16]) < 0), Decimal(0)),
        "trade_fee": sum((_number(row[22]) for row in trades), Decimal(0)),
        "cash_net": sum((_number(row[14]) for row in movements), Decimal(0)),
        "cash_fee": sum((_number(row[14]) for row in movements
                         if str(row[9]).startswith("Комиссия по сделке")
                         or str(row[9]).strip() == "Комиссия"), Decimal(0)),
        "cash_fee_allocated": sum((_number(row[14]) for row in movements
                                   if str(row[9]).startswith("Комиссия по сделке")), Decimal(0)),
        "cash_fee_daily": sum((_number(row[14]) for row in movements
                               if str(row[9]).strip() == "Комиссия"), Decimal(0)),
        "bank_fee_subtotal": cash_subtotal("Списано по тарифам Банка"),
        "cash_start": cash_balance("На начало периода"),
        "cash_end": cash_balance("На конец периода"),
        "quantities": dict(quantities), "positions": positions,
    }


def _parser_totals(path: Path) -> dict:
    report = parse_alfa(path)
    quantities = defaultdict(Decimal)
    for trade in report.trades:
        quantities[trade.isin] += trade.quantity
    return {
        "trade_count": len(report.trades), "cash_row_count": len(report.cash_movements),
        "buy_amount": sum((trade.amount for trade in report.trades if trade.quantity > 0), Decimal(0)),
        "sell_amount": sum((trade.amount for trade in report.trades if trade.quantity < 0), Decimal(0)),
        "trade_fee": sum((trade.commission for trade in report.trades), Decimal(0)),
        "cash_net": sum((row.amount for row in report.cash_movements), Decimal(0)),
        "cash_fee": sum((row.amount for row in report.cash_movements if row.kind in {"trade_fee", "fee"}), Decimal(0)),
        "cash_fee_allocated": sum((row.amount for row in report.cash_movements if row.kind == "trade_fee"), Decimal(0)),
        "cash_fee_daily": sum((row.amount for row in report.cash_movements if row.kind == "fee"), Decimal(0)),
        "bank_fee_subtotal": sum((row.amount for row in report.cash_movements if row.kind in {"trade_fee", "fee"}), Decimal(0)),
        "cash_start": sum((row.start_amount for row in report.cash_snapshots), Decimal(0)),
        "cash_end": sum((row.end_amount for row in report.cash_snapshots), Decimal(0)),
        "quantities": dict(quantities),
        "positions": {row.isin: [row.start_quantity, row.end_quantity] for row in report.positions},
    }


def _compare(actual: dict, expected: dict, description: str) -> None:
    for key, value in expected.items():
        observed = actual[key]
        if isinstance(value, dict):
            observed = {name: str(number) if not isinstance(number, list) else [str(item) for item in number]
                        for name, number in observed.items()}
        elif not isinstance(value, int):
            observed = str(observed)
        if observed != value:
            raise ValueError(f"{description}: {key} differs: {observed!r} != {value!r}")


def check(registry_path: Path, controls_path: Path) -> int:
    registry = json.loads(registry_path.read_text())
    controls = json.loads(controls_path.read_text())
    if controls.get("schema") != 1:
        raise ValueError("Unsupported control schema")
    by_sha = {item["sha256"]: item for item in registry["reports"] if item["broker"] == "alfa"}
    for case in controls["cases"]:
        item = by_sha[case["sha256"]]
        path = Path(item["path"])
        if path.suffix.lower() not in {".xls", ".xlsx"}:
            raise ValueError("This aggregate control supports XLS and XLSX sources only")
        with path.open("rb") as source:
            if hashlib.file_digest(source, "sha256").hexdigest() != case["sha256"]:
                raise ValueError("Registered source changed")
        source_totals = _source_totals(path)
        parser_totals = _parser_totals(path)
        _compare(source_totals, case["expected"], "source")
        _compare(parser_totals, case["expected"], "parser")
        if source_totals != parser_totals:
            raise ValueError("Source and parser totals differ")
        if source_totals["cash_start"] + source_totals["cash_net"] != source_totals["cash_end"]:
            raise ValueError("Cash movement does not reconcile")
        if source_totals["trade_fee"] != -source_totals["cash_fee_allocated"]:
            raise ValueError("Trade and allocated cash commission totals differ")
        if source_totals["cash_fee"] != source_totals["cash_fee_allocated"] + source_totals["cash_fee_daily"]:
            raise ValueError("Cash commission components differ")
        if source_totals["bank_fee_subtotal"] != source_totals["cash_fee"]:
            raise ValueError("Bank fee subtotal differs from cash fee lines")
    return len(controls["cases"])


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("usage: check_alfa_aggregate_controls.py REGISTRY PRIVATE_CONTROLS")
    print(f"Checked {check(Path(sys.argv[1]), Path(sys.argv[2]))} Alfa aggregate control cases")
