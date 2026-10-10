"""Compare parsed numbers with independently read private report cells.

This checks transcription, not the economic meaning of a broker row.
SPDX-License-Identifier: GPL-3.0-or-later
"""

from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import hashlib
import json
import os
from pathlib import Path
import re
import sys
import tempfile

from bs4 import BeautifulSoup

from svoibumagi.alfa import parse_alfa
from svoibumagi.sber import parse_sber


DAY = re.compile(r"\d{2}\.\d{2}\.\d{4}")


def number(value: object, *, empty_zero: bool = False) -> Decimal:
    if value is None or str(value).strip() == "":
        if empty_zero:
            return Decimal(0)
        raise ValueError("empty numeric source cell")
    clean = str(value).replace("\u00a0", "").replace(" ", "").replace(",", ".")
    try:
        result = Decimal(clean)
    except InvalidOperation as exc:
        raise ValueError(f"invalid numeric source cell: {value!r}") from exc
    if not result.is_finite():
        raise ValueError("non-finite numeric source cell")
    return result


def day(value: object) -> str:
    match = DAY.search(str(value))
    if not match:
        raise ValueError("missing source date")
    return datetime.strptime(match.group(), "%d.%m.%Y").date().isoformat()


def _same(actual: object, expected: object, where: str) -> None:
    if actual != expected:
        raise ValueError(f"Source value differs at {where}: parsed={actual!r}, source={expected!r}")


def _sber_rows(path: Path) -> dict[str, list[str]]:
    rows: dict[str, list[str]] = {}
    for table_index, table in enumerate(BeautifulSoup(path.read_bytes(), "html.parser").find_all("table")):
        for row_index, tr in enumerate(table.find_all("tr"), 1):
            rows[f"table:{table_index}:{row_index}"] = [cell.get_text(" ", strip=True).replace("\u00a0", " ")
                                                          for cell in tr.find_all(["td", "th"], recursive=False)]
    return rows


def _alfa_rows(path: Path) -> tuple[dict[str, list[list[object]]], int]:
    if path.suffix.lower() == ".xls":
        import xlrd

        book = xlrd.open_workbook(path, logfile=None)
        return {sheet.name: [sheet.row_values(i) for i in range(sheet.nrows)] for sheet in book.sheets()}, book.datemode
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    return {sheet.title: [list(row) for row in sheet.iter_rows(values_only=True)] for sheet in book}, 0


def _alfa_day(value: object, datemode: int) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (float, int)):
        import xlrd

        return xlrd.xldate_as_datetime(value, datemode).date().isoformat()
    return day(value)


def _loc(item) -> str:
    return f"{item.location.sheet}:{item.location.row}"


def check_sber(path: Path, report) -> Counter[str]:
    rows = _sber_rows(path)
    counts: Counter[str] = Counter()
    for position in report.positions:
        where = _loc(position)
        raw = rows[where]
        _same(position.isin, raw[1].strip(), where + " ISIN")
        _same(position.start_quantity, number(raw[3], empty_zero=True), where + " start quantity")
        _same(position.end_quantity, number(raw[8], empty_zero=True), where + " end quantity")
        counts["position"] += 1
    for cash in report.cash_snapshots:
        where = _loc(cash)
        raw = rows[where]
        _same(cash.start_amount, number(raw[3], empty_zero=True), where + " start cash")
        _same(cash.end_amount, number(raw[5], empty_zero=True), where + " end cash")
        counts["cash_snapshot"] += 1
    movements_by_row: dict[str, list] = defaultdict(list)
    for movement in report.cash_movements:
        movements_by_row[_loc(movement)].append(movement)
    for where, movements in movements_by_row.items():
        raw = rows[where]
        _same(sum((movement.amount for movement in movements), Decimal(0)),
              number(raw[4], empty_zero=True) - number(raw[5], empty_zero=True), where + " net cash movement")
        for movement in movements:
            _same(movement.date.isoformat(), day(raw[0]), where + " cash date")
        counts["cash_movement"] += 1
        counts["cash_component"] += len(movements)
    for trade in report.trades + report.summary_trades:
        where = _loc(trade)
        raw = rows[where]
        sign = 1 if raw[6] == "Покупка" else -1 if raw[6] == "Продажа" else None
        if sign is None:
            raise ValueError(f"Unknown source trade side at {where}")
        _same(trade.quantity, number(raw[7]) * sign, where + " trade quantity")
        _same(trade.price, number(raw[8]), where + " trade price")
        _same(trade.amount, number(raw[9]), where + " trade amount")
        _same(trade.accrued_interest, number(raw[10], empty_zero=True), where + " accrued interest")
        _same(trade.commission, number(raw[11], empty_zero=True) + number(raw[12], empty_zero=True), where + " commission")
        _same(trade.traded_at.isoformat(), day(raw[0]), where + " trade date")
        _same(trade.settled_at.isoformat(), day(raw[1]), where + " settlement date")
        counts["trade"] += 1
    return counts


def check_alfa(path: Path, report) -> Counter[str]:
    sheets, datemode = _alfa_rows(path)
    counts: Counter[str] = Counter()

    def row(item) -> list[object]:
        return sheets[item.location.sheet][item.location.row - 1]

    for position in report.positions:
        where = _loc(position)
        raw = row(position)
        if position.isin:
            next_raw = sheets[position.location.sheet][position.location.row]
            _same(position.isin, str(next_raw[6]).strip(), where + " ISIN")
        _same(position.start_quantity, number(raw[12]), where + " start quantity")
        _same(position.end_quantity, number(raw[15]), where + " end quantity")
        counts["position"] += 1
    for cash in report.cash_snapshots:
        where = _loc(cash)
        raw = row(cash)
        _same(cash.start_amount, number(raw[12]), where + " start cash")
        _same(cash.end_amount, number(raw[15]), where + " end cash")
        counts["cash_snapshot"] += 1
    commission_cols: list[int] = []
    if report.trades:
        trade_rows = sheets["Завершенные сделки"]
        headers = [row for row in trade_rows if any("№ сделки" in str(value) for value in row)
                   and any("ISIN" in str(value) for value in row)]
        if len(headers) != 1:
            raise ValueError("Alfa source trade header is missing or ambiguous")
        commission_cols = [i for i, value in enumerate(headers[0]) if "Комиссия" in str(value) and "банка" in str(value)]
        if len(commission_cols) != 1:
            raise ValueError("Alfa source commission column is missing or ambiguous")
    for trade in report.trades:
        where = _loc(trade)
        raw = row(trade)
        _same(trade.isin, str(raw[12]).strip(), where + " ISIN")
        _same(trade.quantity, number(raw[16]), where + " trade quantity")
        _same(trade.price, number(raw[18]), where + " trade price")
        _same(trade.amount, number(raw[19]), where + " trade amount")
        _same(trade.commission, number(raw[commission_cols[0]], empty_zero=True), where + " commission")
        _same(trade.traded_at.isoformat(), _alfa_day(raw[7], datemode), where + " trade date")
        _same(trade.settled_at.isoformat(), _alfa_day(raw[10], datemode), where + " settlement date")
        counts["trade"] += 1
    by_row: dict[str, list] = defaultdict(list)
    for movement in report.cash_movements:
        by_row[_loc(movement)].append(movement)
    for where, movements in by_row.items():
        raw = row(movements[0])
        _same(sum((movement.amount for movement in movements), Decimal(0)), number(raw[14]),
              where + " net cash movement")
        label = str(raw[9]).strip()
        expected_kinds = (
            ["trade_cash"] if label.startswith("Расчеты по сделке ") else
            ["trade_fee"] if label.startswith("Комиссия по сделке ") else
            ["accrued_interest"] if label == "НКД по сделке" or label.startswith("НКД по сделке ") else
            ["fee"] if label == "Комиссия" else
            ["tax_withheld"] if label == "НДФЛ" else None
        )
        if label == "Перевод":
            comment = str(raw[10]).strip().casefold()
            expected_kinds = (
                ["deposit"] if comment.startswith(('из ао "альфа-банк"', "с субсчета ")) else
                ["withdrawal"] if comment.startswith("возврат средств по дог.") else
                ["coupon"] if comment.startswith("погашение купона ") else
                ["dividend", "tax_withheld"] if comment.startswith("дивиденды ") else None
            )
        if expected_kinds is None:
            raise ValueError(f"Unknown source cash category at {where}: {label}")
        _same([movement.kind for movement in movements], expected_kinds, where + " cash category")
        counts["cash_semantic_source_row"] += 1
        source_rows = sheets[movements[0].location.sheet]
        source_day = next((_alfa_day(previous[2], datemode) for previous in
                           reversed(source_rows[:movements[0].location.row])
                           if len(previous) > 2 and previous[2] not in (None, "")
                           and (isinstance(previous[2], (date, datetime, float, int)) or DAY.search(str(previous[2])))), None)
        if source_day is None:
            raise ValueError(f"No source cash date at {where}")
        for movement in movements:
            _same(movement.date.isoformat(), source_day, where + " cash date")
        counts["cash_source_row"] += 1
        counts["cash_component"] += len(movements)
    for subtotal in report.cash_flow_subtotals:
        where = _loc(subtotal)
        _same(subtotal.amount, number(row(subtotal)[14]), where + " subtotal")
        counts["cash_subtotal"] += 1
        if subtotal.kind == "bank_fee":
            source_total = Decimal(0)
            for row_number, raw in enumerate(sheets[" Движение ДС"], 1):
                label = str(raw[9]).strip() if len(raw) > 9 else ""
                if label == "Комиссия":
                    expected_kind = "fee"
                elif label.startswith("Комиссия по сделке "):
                    expected_kind = "trade_fee"
                else:
                    continue
                source_total += number(raw[14])
                movements = by_row.get(f" Движение ДС:{row_number}", [])
                _same([movement.kind for movement in movements], [expected_kind],
                      f" Движение ДС:{row_number} fee category")
                counts["bank_fee_source_row"] += 1
            _same(subtotal.amount, source_total, where + " bank fee source total")
            counts["bank_fee_subtotal"] += 1
    return counts


def audit(registry_path: Path) -> dict:
    registry_bytes = registry_path.read_bytes()
    registry = json.loads(registry_bytes)
    result: Counter[str] = Counter()
    for item in registry["reports"]:
        path = Path(item["path"])
        with path.open("rb") as source:
            _same(hashlib.file_digest(source, "sha256").hexdigest(), item["sha256"], "registered report hash")
        report = (parse_sber if item["broker"] == "sber" else parse_alfa)(path)
        try:
            counts = (check_sber if item["broker"] == "sber" else check_alfa)(path, report)
        except (IndexError, KeyError, ValueError) as exc:
            raise ValueError(f"Source value audit failed for {item['sha256'][:12]}: {exc}") from exc
        result[f"{item['broker']}_reports"] += 1
        for key, value in counts.items():
            result[f"{item['broker']}_{key}"] += value
    return {"schema": 1, "registry_sha256": hashlib.sha256(registry_bytes).hexdigest(),
            "reports": len(registry["reports"]), "counts": dict(sorted(result.items()))}


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("usage: audit_report_values.py PRIVATE_REGISTRY PRIVATE_OUTPUT")
    result = audit(Path(sys.argv[1]))
    output = Path(sys.argv[2])
    output.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    with tempfile.NamedTemporaryFile("w", dir=output.parent, prefix=".value-audit-", suffix=".tmp",
                                     delete=False, encoding="utf-8") as stream:
        temporary = Path(stream.name)
        json.dump(result, stream, ensure_ascii=False, indent=2)
        stream.write("\n")
    try:
        os.chmod(temporary, 0o600)
        temporary.replace(output)
    finally:
        temporary.unlink(missing_ok=True)
    print(f"Checked source values in {result['reports']} reports: {result['counts']}")


if __name__ == "__main__":
    main()
