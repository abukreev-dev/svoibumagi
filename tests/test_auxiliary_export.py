"""Synthetic auxiliary Sber export checks. SPDX-License-Identifier: GPL-3.0-or-later"""

from datetime import datetime

from openpyxl import Workbook
import pytest

from svoibumagi.auxiliary_export import check


def _table(rows):
    return "<table>" + "".join("<tr>" + "".join(f"<td>{cell}</td>" for cell in row) + "</tr>" for row in rows) + "</table>"


def test_auxiliary_export_matches_broker_report_and_detects_changed_amount(tmp_path):
    position = ["Synthetic", "RU000A000001", "RUB", "0"] + ["0"] * 4 + ["1"] + ["0"] * 9
    html = "Отчет брокера за период с 29.01.2025 по 29.01.2025 Договор SYN123 от 01.01.2025 " + "".join([
        _table([["Основной рынок", "Начало периода", "Конец периода"],
                ["Наименование", "ISIN ценной бумаги"] + ["-"] * 16, position]),
        _table([["Торговая площадка", "Валюта", "Конец периода"], ["1", "2", "3"],
                ["Торговый счет, Основной рынок", "RUB", "1", "100", "-20", "80", "0", "0", "80"]]),
        _table([["Дата", "Торговая площадка", "Описание операции", "Валюта", "Сумма зачисления", "Сумма списания"],
                ["1", "2", "3", "4", "5", "6"],
                ["29.01.2025", "Основной рынок", "Вывод д/с", "RUB", "", "20"]]),
        _table([["Дата заключения", "Дата расчетов", "Время заключения", "Наименование ЦБ", "Код ЦБ", "Валюта", "Вид", "Количество, шт.", "Цена", "Сумма", "НКД", "Комиссия Брокера", "Комиссия Биржи", "Номер сделки", "Комментарий", "Статус сделки"],
                list(map(str, range(1, 17))),
                ["28.01.2025", "29.01.2025", "12:00", "Synthetic", "SYN", "RUB", "Покупка", "1", "10", "10", "0", "0", "0", "deal-1", "", "ЗИ"]]),
        _table([["Наименование", "Код", "ISIN ценной бумаги"], ["1", "2", "3"],
                ["Synthetic", "SYN", "RU000A000001"]]),
    ])
    report_path = tmp_path / "report.html"
    report_path.write_text(html)
    workbook = Workbook()
    trades = workbook.active
    trades.title = "Сделки"
    trades.append(["Номер договора", "Номер сделки", "Дата заключения"] + [""] * 9)
    trades.append(["SYN123", "deal-1", datetime(2025, 1, 28, 12), "SYN", "Облигация", "Фондовый рынок",
                   "Покупка", 1, 100, None, 10, "RUB"])
    cash = workbook.create_sheet("Движение ДС")
    cash.append(["Номер договора", "Дата подачи поручения", "Дата исполнения поручения", "Операция"] + [""] * 9)
    cash.append(["SYN123", datetime(2025, 1, 29, 8), datetime(2025, 1, 29, 9), "Вывод со счёта",
                 None, 20, "RUB", None, None, None, "Исполнена"])
    export_path = tmp_path / "operations.xlsx"
    workbook.save(export_path)
    assert check(export_path, report_path) == {"matched_trades": 1, "matched_withdrawals": 1,
                                               "bond_price_scale_ten": 1}
    trades.cell(2, 11, 11)
    workbook.save(export_path)
    with pytest.raises(ValueError, match="economics differ"):
        check(export_path, report_path)
