"""Synthetic broker-report controls. SPDX-License-Identifier: GPL-3.0-or-later"""

from decimal import Decimal as D
from datetime import date
from dataclasses import replace
import json
import zipfile

from openpyxl import Workbook
import pytest

from svoibumagi.alfa import UnsupportedReport, _num, _period, _sheet_rows, parse_alfa
from svoibumagi.backup import create_backup, restore_backup
from svoibumagi.metrics import recorded_income
from svoibumagi.reconcile import reconcile
from svoibumagi.sber import parse_sber
from svoibumagi.storage import Storage
from scripts.audit_report_values import check_alfa
from tests.source_fixture import source_row


def _table(rows):
    return "<table>" + "".join("<tr>" + "".join(f"<td>{cell}</td>" for cell in row) + "</tr>" for row in rows) + "</table>"


def test_broker_adapters_reject_reversed_report_period(tmp_path):
    path = tmp_path / "reversed.html"
    path.write_text("<html><body>Отчет брокера за период с 31.01.2025 по 01.01.2025 "
                    "Договор SYN123 от 01.01.2025</body></html>")
    with pytest.raises(UnsupportedReport, match="ends before it starts"):
        parse_sber(path)
    store = Storage(tmp_path / "private")
    account = store.create_account("sber", "brokerage", "Synthetic")
    with pytest.raises(UnsupportedReport, match="ends before it starts"):
        store.stage_file(account, path)
    assert store.list_attempts() == []
    assert list(store.files.iterdir()) == []
    with pytest.raises(UnsupportedReport, match="ends before it starts"):
        _period({"Динамика позиций": [["31.01.2025 - 01.01.2025"]]})


def test_xlsx_expansion_and_cell_limits_are_checked_before_reading_all_rows(tmp_path, monkeypatch):
    import svoibumagi.alfa as alfa

    compressed = tmp_path / "expanded.xlsx"
    with zipfile.ZipFile(compressed, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", b"x" * 2_000)
    monkeypatch.setattr(alfa, "MAX_XLSX_EXPANDED_BYTES", 1_000)
    with pytest.raises(UnsupportedReport, match="expansion limit"):
        _sheet_rows(compressed)
    store = Storage(tmp_path / "private")
    account = store.create_account("alfa", "brokerage", "Synthetic")
    with pytest.raises(UnsupportedReport, match="expansion limit"):
        store.stage_file(account, compressed)
    assert store.list_attempts() == [] and list(store.files.iterdir()) == []

    book = Workbook()
    sheet = book.active
    for row in range(1, 4):
        for column in range(1, 4):
            sheet.cell(row, column, row * column)
    path = tmp_path / "cells.xlsx"
    book.save(path)
    monkeypatch.setattr(alfa, "MAX_XLSX_EXPANDED_BYTES", 128_000_000)
    monkeypatch.setattr(alfa, "MAX_WORKBOOK_CELLS", 8)
    with pytest.raises(UnsupportedReport, match="cell limit"):
        _sheet_rows(path)


def test_report_numbers_reject_nonfinite_values(tmp_path):
    for value in ("NaN", "Infinity", "-Infinity", float("nan"), float("inf")):
        with pytest.raises(ValueError, match="finite"):
            _num(value)

    html = "Отчет брокера за период с 01.01.2025 по 31.01.2025 Договор SYN123 от 01.01.2025 " + "".join([
        _table([["Основной рынок", "Начало периода", "Конец периода"],
                ["Наименование", "ISIN ценной бумаги"] + ["-"] * 16]),
        _table([["Торговая площадка", "Валюта", "Курс на конец периода", "Начало периода", "Изменение за период", "Конец периода"],
                ["1", "2", "3", "4", "5", "6"],
                ["Торговый счет, Основной рынок", "RUB", "1", "10", "0", "NaN", "0", "0", "NaN"]]),
    ])
    path = tmp_path / "nonfinite.html"
    path.write_text(html)
    report = parse_sber(path)
    assert not report.cash_snapshots
    assert any("invalid cash snapshot" in item.reason for item in report.unparsed)


def test_sber_dividend_gross_and_tax_share_one_source_row(tmp_path):
    def source(description):
        return "Отчет брокера за период с 01.01.2025 по 31.01.2025 Договор SYN123 от 01.01.2025 " + "".join([
            _table([["Основной рынок", "Начало периода", "Конец периода"],
                    ["Наименование", "ISIN ценной бумаги"] + ["-"] * 16]),
            _table([["Торговая площадка", "Валюта", "Курс на конец периода", "Начало периода", "Изменение за период", "Конец периода"],
                    ["1", "2", "3", "4", "5", "6"],
                    ["Торговый счет, Основной рынок", "RUB", "1", "10", "87", "97", "0", "0", "97"]]),
            _table([["Дата", "Торговая площадка", "Описание операции", "Валюта", "Сумма зачисления", "Сумма списания"],
                    ["1", "2", "3", "4", "5", "6"],
                    ["15.01.2025", "Основной рынок", description, "RUB", "87", "0"]]),
        ])

    report_path = tmp_path / "dividend.html"
    report_path.write_text(source("Дивиденды Synthetic; Налог удержан Дополнительная информация: Дивиденды 100.00 RUR по курсу ЦБ 1"))
    report = parse_sber(report_path)
    assert report.variant == "sber-html-market-v7"
    assert [(row.kind, row.amount, row.withheld_tax, row.component) for row in report.cash_movements] == [
        ("dividend", D("100"), D("13"), ""), ("tax_withheld", D("-13"), None, "dividend_tax")]
    assert not report.unparsed
    assert next(check for check in reconcile(report) if check.kind == "cash").status == "matched"
    store = Storage(tmp_path / "private")
    account = store.create_account("sber", "brokerage", "Synthetic")
    attempt, _ = store.stage_file(account, report_path)
    review = store.review_attempt(attempt)
    decisions = {key: {"reason": "synthetic missing position snapshot"} for key in review["required_decisions"]}
    with store._connect() as con:
        staged_row = con.execute(
            "SELECT sheet,row_number,raw_json,classification,warning FROM source_rows "
            "WHERE attempt_id=? AND classification='cash_movement'", (attempt,)
        ).fetchone()
        con.execute(
            "DELETE FROM source_rows WHERE attempt_id=? AND sheet=? AND row_number=? AND classification=?",
            (attempt, staged_row[0], staged_row[1], staged_row[3]),
        )
    with pytest.raises(ValueError, match="matching staged source row"):
        store.confirm_attempt(attempt, decisions)
    assert store.list_events(account) == []
    with store._connect() as con:
        assert con.execute("SELECT state FROM import_attempts WHERE id=?", (attempt,)).fetchone()[0] == "awaiting_review"
    with store._connect() as con:
        con.execute(
            "INSERT INTO source_rows(attempt_id,sheet,row_number,raw_json,classification,warning) "
            "VALUES(?,?,?,?,?,?)", (attempt, *staged_row),
        )
    assert store.confirm_attempt(attempt, decisions)["inserted"] == 2
    assert {(event["sheet"], event["row_number"], event["payload"]["kind"])
            for event in store.list_events(account)} == {
                ("table:2", 3, "dividend"), ("table:2", 3, "tax_withheld")}
    with store._connect() as con:
        assert {row[0] for row in con.execute("SELECT component FROM event_observations")} == {"", "dividend_tax"}

    legacy = Storage(tmp_path / "legacy-private")
    legacy_account = legacy.create_account("sber", "brokerage", "Synthetic legacy")
    legacy_attempt, _ = legacy.stage_file(legacy_account, report_path)
    with legacy._connect() as con:
        source_id = con.execute("SELECT source_id FROM import_attempts WHERE id=?", (legacy_attempt,)).fetchone()[0]
        # Reproduce a pre-v17 row that was written before its import was confirmed.
        con.execute("DROP TRIGGER validate_ledger_source_row")
        con.execute("DROP TRIGGER validate_observation_source_row")
        old_event = con.execute(
            "INSERT INTO ledger_events(account_id,source_id,sheet,row_number,event_type,broker_event_id,fingerprint,payload_json,created_at) "
            "VALUES(?,?,?,?,?,?,?,?,?)", (legacy_account, source_id, "table:2", 3, "cash", None, "old-net",
                                          json.dumps({"kind": "dividend", "date": "2025-01-15", "amount": "87",
                                                      "currency": "RUB"}), "2025-01-15"),
        ).lastrowid
        con.execute("INSERT INTO event_observations(event_id,source_id,sheet,row_number,event_type) VALUES(?,?,?,?,?)",
                    (old_event, source_id, "table:2", 3, "cash"))
    legacy_review = legacy.review_attempt(legacy_attempt)
    legacy_decisions = {key: {"reason": "synthetic legacy source reviewed"}
                        for key in legacy_review["required_decisions"]}
    with pytest.raises(ValueError, match="Previously imported source cells"):
        legacy.confirm_attempt(legacy_attempt, legacy_decisions)
    assert len(legacy.list_events(legacy_account)) == 1

    overlapping = tmp_path / "dividend-overlap.html"
    overlapping.write_text(report_path.read_text() + "<!-- another export -->")
    overlap_attempt, _ = legacy.stage_file(legacy_account, overlapping)
    overlap_review = legacy.review_attempt(overlap_attempt)
    assert any(row["reason"] == "source cells already have a different cash event"
               for row in overlap_review["ambiguous_events"])
    overlap_decisions = {key: {"reason": "synthetic legacy source reviewed"}
                         for key in overlap_review["required_decisions"]}
    with pytest.raises(ValueError, match="different cash event"):
        legacy.confirm_attempt(overlap_attempt, overlap_decisions)
    assert len(legacy.list_events(legacy_account)) == 1

    fallback = tmp_path / "net-only.html"
    for description in ("Выплата дивидендов. Налог удержан.",
                        "Дивиденды Synthetic; Налог удержан Дополнительная информация: Дивиденды 80.00 RUR по курсу ЦБ 1"):
        fallback.write_text(source(description))
        net = parse_sber(fallback)
        assert [(row.kind, row.amount) for row in net.cash_movements] == [("dividend_net", D("87"))]
        assert any("withholding amount not separately verified" in row.reason for row in net.unparsed)


def test_sber_coupon_credit_and_tax_debit_are_not_external_transfers(tmp_path):
    html = "Отчет брокера за период с 01.01.2025 по 31.01.2025 Договор SYN123 от 01.01.2025 " + "".join([
        _table([["Основной рынок", "Начало периода", "Конец периода"],
                ["Наименование", "ISIN ценной бумаги"] + ["-"] * 16]),
        _table([["Торговая площадка", "Валюта", "Курс на конец периода", "Начало периода", "Изменение за период", "Конец периода"],
                ["1", "2", "3", "4", "5", "6"],
                ["Торговый счет, Основной рынок", "RUB", "1", "10", "87", "97", "0", "0", "97"]]),
        _table([["Дата", "Торговая площадка", "Описание операции", "Валюта", "Сумма зачисления", "Сумма списания"],
                ["1", "2", "3", "4", "5", "6"],
                ["15.01.2025", "Основной рынок", "Зачисление д/с (купон 3 по ОФЗ 26248)", "RUB", "100", "0"],
                ["16.01.2025", "Основной рынок", "Списание д/с. Налог на доходы физ.лиц", "RUB", "0", "13"]]),
    ])
    path = tmp_path / "coupon-and-tax.html"
    path.write_text(html)
    report = parse_sber(path)
    assert [(row.kind, row.amount) for row in report.cash_movements] == [
        ("coupon", D("100")), ("tax_withheld", D("-13"))]
    assert next(check for check in reconcile(report) if check.kind == "cash").status == "matched"
    store = Storage(tmp_path / "private")
    account = store.create_account("sber", "brokerage", "Synthetic")
    attempt, _ = store.stage_file(account, path)
    review = store.review_attempt(attempt)
    decisions = {key: {"reason": "synthetic missing position snapshot"} for key in review["required_decisions"]}
    assert store.confirm_attempt(attempt, decisions)["inserted"] == 2
    events = store.list_events(account)
    assert {event["payload"]["kind"] for event in events} == {"coupon", "tax_withheld"}
    income = recorded_income(events, date(2025, 1, 31))
    assert income["coupon"] == D("100")
    assert income["tax_withheld"] == D("13")

    legacy = Storage(tmp_path / "legacy-private")
    legacy_account = legacy.create_account("sber", "brokerage", "Synthetic legacy")
    old_attempt, _ = legacy.stage_file(legacy_account, path)
    with legacy._connect() as con:
        source_id = con.execute("SELECT source_id FROM import_attempts WHERE id=?", (old_attempt,)).fetchone()[0]
        con.execute("DROP TRIGGER validate_ledger_source_row")
        con.execute("DROP TRIGGER validate_observation_source_row")
        old_event = con.execute(
            "INSERT INTO ledger_events(account_id,source_id,sheet,row_number,event_type,broker_event_id,fingerprint,payload_json,created_at) "
            "VALUES(?,?,?,?,?,?,?,?,?)",
            (legacy_account, source_id, "table:2", 3, "cash", None, "old-coupon-as-deposit",
             json.dumps({"kind": "deposit", "date": "2025-01-15", "amount": "100", "currency": "RUB"}),
             "2025-01-15"),
        ).lastrowid
        con.execute("INSERT INTO event_observations(event_id,source_id,sheet,row_number,event_type) VALUES(?,?,?,?,?)",
                    (old_event, source_id, "table:2", 3, "cash"))
        # Emulate an import made before stored source rows were protected.
        con.execute("DROP TRIGGER protect_last_referenced_source_row_delete")
        con.execute("DELETE FROM source_rows WHERE attempt_id=?", (old_attempt,))
    overlapping = tmp_path / "coupon-overlap.html"
    overlapping.write_text(html + "<!-- separate export of the same source cells -->")
    new_attempt, _ = legacy.stage_file(legacy_account, overlapping)
    review = legacy.review_attempt(new_attempt)
    assert any(row["reason"] == "source cells already have a different cash event"
               for row in review["ambiguous_events"])
    decisions = {key: {"reason": "synthetic source review"} for key in review["required_decisions"]}
    with pytest.raises(ValueError, match="different cash event"):
        legacy.confirm_attempt(new_attempt, decisions)
    assert len(legacy.list_events(legacy_account)) == 1


def test_sber_structural_rows_are_preserved_without_ledger_events(tmp_path):
    html = "Отчет брокера за период с 01.01.2025 по 31.01.2025 Договор SYN123 от 01.01.2025 " + "".join([
        _table([["Основной рынок", "Начало периода", "Конец периода"], ["Наименование", "ISIN ценной бумаги"],
                list(map(str, range(1, 19))), ["Площадка: Фондовый рынок"]]),
        _table([["Торговая площадка", "Валюта", "Курс на конец периода", "Начало периода", "Изменение за период", "Конец периода"],
                ["Счёт", "Валюта"], list(map(str, range(1, 11))),
                ["Основной рынок"] + list(map(str, range(2, 11))),
                ["Торговый счет, Основной рынок", "USD", "1", "0", "0", "0", "0", "0", "0"]]),
    ])
    path = tmp_path / "structural.html"
    path.write_text(html)
    report = parse_sber(path)
    assert len(report.excluded) == 5
    assert {row.reason for row in report.excluded} == {"structural Sber row; no ledger event"}
    assert not report.unparsed
    assert not report.positions and not report.cash_snapshots
    path.write_text(html.replace("</table>", "<tr><td>Synthetic unknown row</td></tr></table>", 1))
    assert any(row.reason == "unrecognized nonempty Sber source row"
               for row in parse_sber(path).unparsed)


def test_sber_broker_subtotal_difference_is_visible(tmp_path):
    html = "Отчет брокера за период с 01.01.2025 по 31.01.2025 Договор SYN123 от 01.01.2025 " + "".join([
        _table([["Основной рынок", "Начало периода", "Конец периода"], ["Наименование", "ISIN ценной бумаги"]]),
        _table([["Торговая площадка", "Валюта", "Конец периода"], ["Счёт", "Валюта"],
                ["Торговый счет, Основной рынок", "RUB", "1", "50", "-20", "30", "0", "0", "30"]]),
        _table([["Дата", "Описание операции"], ["Дата", "Описание"],
                ["15.01.2025", "Основной рынок", "Сделка", "RUB", "", "20"],
                ["Итого, RUB", "0", "20"]]),
    ])
    path = tmp_path / "total.html"
    path.write_text(html)
    report = parse_sber(path)
    assert len(report.source_totals) == 2
    assert sum(check.kind == "source_total" and check.status == "matched" for check in reconcile(report)) == 2
    path.write_text(html.replace("Итого, RUB</td><td>0</td><td>20", "Итого, RUB</td><td>0</td><td>21"))
    changed = parse_sber(path)
    assert any(check.kind == "source_total" and check.status == "difference" for check in reconcile(changed))
    store = Storage(tmp_path / "subtotal-private")
    account = store.create_account("sber", "brokerage", "Synthetic subtotal")
    attempt, _ = store.stage_file(account, path)
    assert any(key.startswith("check:source_total:") for key in store.review_attempt(attempt)["required_decisions"])
    with pytest.raises(ValueError):
        store.confirm_attempt(attempt)
    assert store.list_events(account) == []


def test_sber_overview_cross_checks_position_and_cash_totals(tmp_path):
    html = "Отчет брокера за период с 01.01.2025 по 31.01.2025 Договор SYN123 от 01.01.2025 " + "".join([
        _table([["Торговая площадка", "Начало периода", "Конец периода", "Изменение за период"],
                ["Ценные бумаги", "Денежные средства"],
                ["Основной рынок", "0", "50", "50", "0", "30", "30", "0", "-20", "-20"],
                ["Итого", "50", "30", "-20"]]),
        _table([["Основной рынок", "Начало периода", "Конец периода"], ["Наименование", "ISIN ценной бумаги"],
                ["Synthetic", "RU000A000001", "RUB"] + ["0"] * 15,
                ["Итого по Основному рынку, RUB", "0", "0", "", "", "", "0", "0", "", "0", "", "", ""]]),
        _table([["Торговая площадка", "Валюта", "Конец периода"], ["Счёт", "Валюта"],
                ["Торговый счет, Основной рынок", "RUB", "1", "50", "-20", "30", "0", "0", "30"],
                ["Итого по площадке Основной рынок", "50", "-20", "30", "0", "0", "30"]]),
        _table([["Дата", "Описание операции"], ["Дата", "Описание"],
                ["15.01.2025", "Основной рынок", "Сделка", "RUB", "", "20"]]),
    ])
    path = tmp_path / "overview.html"
    path.write_text(html)
    report = parse_sber(path)
    overview = [row for row in report.source_totals if row.location.sheet == "table:0"]
    assert len(overview) == 12
    assert all(check.status == "matched" for check in reconcile(report) if check.kind == "source_total")
    path.write_text(html.replace("Итого</td><td>50</td><td>30", "Итого</td><td>50</td><td>31"))
    changed = parse_sber(path)
    assert any(check.kind == "source_total" and check.status == "difference" for check in reconcile(changed))


def test_sber_html_report_reconciles(tmp_path):
    position_head = ["Наименование", "ISIN ценной бумаги", "Валюта рыночной цены", "Количество, шт"] + ["-"] * 4 + ["Количество, шт"] + ["-"] * 9
    position = ["Synthetic", "RU000A000001", "RUB", "1"] + ["0"] * 4 + ["3"] + ["0"] * 9
    html = "Отчет брокера за период с 01.01.2025 по 31.01.2025 Договор SYN123 от 01.01.2025 " + "".join([
        _table([["Основной рынок", "Начало периода", "Конец периода"], position_head, position]),
        _table([["Торговая площадка", "Валюта", "Курс на конец периода", "Начало периода", "Изменение за период", "Конец периода"],
                ["1", "2", "3", "4", "5", "6"], ["Торговый счет, Основной рынок", "RUB", "1", "50", "-20", "30", "0", "0", "30"]]),
        _table([["Дата", "Торговая площадка", "Описание операции", "Валюта", "Сумма зачисления", "Сумма списания"],
                ["1", "2", "3", "4", "5", "6"], ["15.01.2025", "Основной рынок", "Сделка от 15.01.2025", "RUB", "", "20"]]),
        _table([["Дата заключения", "Дата расчетов", "Время заключения", "Наименование ЦБ", "Код ЦБ", "Валюта", "Вид", "Количество, шт.", "Цена", "Сумма", "НКД", "Комиссия Брокера", "Комиссия Биржи", "Номер сделки", "Комментарий", "Статус сделки"],
                list(map(str, range(1, 17))), ["15.01.2025", "15.01.2025", "12:00", "Synthetic", "SYN", "RUB", "Покупка", "2", "10", "20", "0", "0", "0", "deal-1", "", "ЗИ"],
                ["31.01.2025", "03.02.2025", "12:00", "Synthetic", "SYN", "RUB", "Покупка", "1", "10", "10", "0", "0", "0", "deal-2", "", "ЗО"]]),
        _table([["Наименование", "Код", "ISIN ценной бумаги"], ["1", "2", "3"], ["Synthetic", "SYN", "RU000A000001"]]),
    ])
    path = tmp_path / "report.html"
    path.write_text(html)
    report = parse_sber(path)
    assert report.account_key is not None
    assert report.account_number == "SYN123"
    assert len(report.trades) == 1
    assert {(row.location.sheet, row.location.row, row.reason) for row in report.excluded} == {
        ("table:4", 3, "instrument reference used for ISIN mapping; no ledger event"),
        ("table:3", 4, "not executed at report end: ЗО"),
    }
    assert len(report.cash_movements) == 1
    assert all(check.status == "matched" for check in reconcile(report))
    old_pending = Storage(tmp_path / "old-pending-private")
    old_account = old_pending.create_account("sber", "brokerage", "Synthetic old preview")
    old_attempt, _ = old_pending.stage_file(old_account, path)
    with old_pending._connect() as con:
        con.execute("DROP TRIGGER protect_import_attempt_fields")
        con.execute("UPDATE import_attempts SET adapter_version='sber-html-market-v2' WHERE id=?",
                    (old_attempt,))
    assert old_pending.review_attempt(old_attempt)["requires_restage"] is True
    with pytest.raises(ValueError, match="Restage broker report"):
        old_pending.confirm_attempt(old_attempt)
    assert old_pending.list_events(old_account) == []
    with_accrued = html.replace("<td>20</td><td>0</td><td>0</td><td>0</td><td>deal-1</td>",
                                "<td>20</td><td>2</td><td>0</td><td>0</td><td>deal-1</td>", 1)
    assert with_accrued != html
    accrued_path = tmp_path / "accrued.html"
    accrued_path.write_text(with_accrued)
    accrued_report = parse_sber(accrued_path)
    assert accrued_report.trades[0].accrued_interest == D("2")
    assert any(check.kind == "settlement" and check.status == "difference"
               for check in reconcile(accrued_report))
    numeric_name = tmp_path / "numeric-name.html"
    numeric_name.write_text(html.replace("<td>Synthetic</td>", "<td>26207</td>", 1))
    numeric_report = parse_sber(numeric_name)
    assert len(numeric_report.positions) == 1 and numeric_report.positions[0].name == "26207"
    assert all(check.status == "matched" for check in reconcile(numeric_report))
    reference_table = _table([["Наименование", "Код", "ISIN ценной бумаги"], ["1", "2", "3"],
                              ["Synthetic", "SYN", "RU000A000001"]])
    conflicting_table = _table([["Наименование", "Код", "ISIN ценной бумаги"], ["1", "2", "3"],
                                ["Synthetic", "SYN", "RU000A000001"],
                                ["Other", "SYN", "RU000A000002"]])
    conflicting_path = tmp_path / "conflicting-reference.html"
    conflicting_path.write_text(html.replace(reference_table, conflicting_table))
    conflicting_report = parse_sber(conflicting_path)
    assert any(row.location.sheet == "table:4" and row.location.row == 4
               and "conflicts with earlier ISIN" in row.reason for row in conflicting_report.unparsed)
    assert conflicting_report.trades[0].isin == "RU000A000001"
    duplicate_section = tmp_path / "duplicate-section.html"
    duplicate_section.write_text(html + _table([
        ["Дата", "Торговая площадка", "Описание операции", "Валюта", "Сумма зачисления", "Сумма списания"],
        ["1", "2", "3", "4", "5", "6"],
    ]))
    with pytest.raises(UnsupportedReport, match="ambiguous"):
        parse_sber(duplicate_section)
    summary_path = tmp_path / "summary.html"
    summary_path.write_text(html.replace("deal-1", ""))
    summary = parse_sber(summary_path)
    assert len(summary.summary_trades) == 1 and not summary.trades
    outside_settlement = tmp_path / "outside-settlement.html"
    outside_settlement.write_text(html.replace(
        "<td>15.01.2025</td><td>15.01.2025</td><td>12:00</td>",
        "<td>31.01.2025</td><td>03.02.2025</td><td>12:00</td>", 1))
    outside = parse_sber(outside_settlement)
    assert not outside.trades
    assert any(row.location.sheet == "table:3" and row.location.row == 3
               and "executed trade settlement outside report period" in row.reason for row in outside.unparsed)
    earlier_execution = tmp_path / "earlier-execution.html"
    earlier_execution.write_text(html.replace(
        "<td>15.01.2025</td><td>15.01.2025</td><td>12:00</td>",
        "<td>31.12.2024</td><td>15.01.2025</td><td>12:00</td>", 1))
    assert len(parse_sber(earlier_execution).trades) == 1
    assert any("without broker ID" in row.reason for row in summary.unparsed)
    assert all(check.status == "matched" for check in reconcile(summary) if check.kind in {"position", "cash"})
    summary_storage = Storage(tmp_path / "summary-private")
    summary_account = summary_storage.create_account("sber", "brokerage", "Synthetic summary")
    summary_attempt, _ = summary_storage.stage_file(summary_account, summary_path)
    summary_review = summary_storage.review_attempt(summary_attempt)
    assert len(summary_review["instrument_references"]) == 1
    assert len(summary_review["reference_rows"]) == 1
    assert len(summary_review["other_excluded"]) == 1
    assert not any((row["sheet"], row["row"]) == ("table:4", 3)
                   for row in summary_review["diagnostic_rows"])
    assert not any(item["key"].startswith("event:trade:") for item in summary_review["event_candidates"])
    with pytest.raises(ValueError, match="Explicit review"):
        summary_storage.confirm_attempt(summary_attempt)
    summary_decisions = {key: {"action": "acknowledge", "reason": "Synthetic source checked"}
                         for key in summary_review["required_decisions"]}
    summary_key = next(key for key in summary_decisions if key.startswith("summary:"))
    with pytest.raises(ValueError, match="control_only or promote"):
        summary_storage.confirm_attempt(summary_attempt, summary_decisions)
    summary_decisions[summary_key]["action"] = "promote"
    assert summary_storage.confirm_attempt(summary_attempt, summary_decisions) == {"inserted": 2, "observed_duplicates": 0}
    promoted = next(event for event in summary_storage.list_events(summary_account) if event["event_type"] == "trade")
    assert promoted["payload"]["aggregate_from_report"] is True
    assert promoted["sheet"] == "table:3" and promoted["row_number"] == 3
    assert summary_storage.committed_reports(summary_account)[0]["unparsed"] == []
    assert all(check["kind"] != "completeness" for check in summary_storage.committed_reports(summary_account)[0]["checks"])
    assert summary_storage.latest_snapshots()[0]["unparsed"] == []
    with summary_storage._connect() as con:
        assert con.execute("SELECT COUNT(*) FROM audit_decisions WHERE action LIKE 'summary:%:promote'").fetchone()[0] == 1
    control_storage = Storage(tmp_path / "control-private")
    control_account = control_storage.create_account("sber", "brokerage", "Synthetic control")
    control_attempt, _ = control_storage.stage_file(control_account, summary_path)
    control_review = control_storage.review_attempt(control_attempt)
    control_decisions = {key: {"action": "control_only" if key.startswith("summary:") else "acknowledge",
                               "reason": "Synthetic control"} for key in control_review["required_decisions"]}
    assert control_storage.confirm_attempt(control_attempt, control_decisions) == {"inserted": 1, "observed_duplicates": 0}
    assert control_storage.committed_reports(control_account)[0]["unparsed"]
    aggregate = control_review["preview"]["summary_trades"][0]
    with control_storage._connect() as con:
        detail_source = con.execute(
            "INSERT INTO source_files(account_id,sha256,original_name,private_path,size,detected_broker,created_at) "
            "VALUES(?,?,?,?,?,?,?)",
            (control_account, "synthetic-detail", "detail.html", "/synthetic", 1, "sber", "2025-01-01"),
        ).lastrowid
        for index in (1, 2):
            source_row(con, detail_source, "synthetic", index, "trade")
            detail = {key: value for key, value in aggregate.items() if key != "location"}
            detail["broker_id"] = f"detail-{index}"
            for key in ("quantity", "amount", "commission"):
                detail[key] = str(D(detail[key]) / 2)
            con.execute(
                "INSERT INTO ledger_events(account_id,source_id,sheet,row_number,event_type,broker_event_id,fingerprint,payload_json,created_at) "
                "VALUES(?,?,?,?,?,?,?,?,?)",
                (control_account, detail_source, "synthetic", index, "trade", detail["broker_id"],
                 f"detail-{index}", json.dumps(detail), "2025-01-01"),
            )
    effective = next(item for item in control_storage.committed_reports(control_account)
                     if item["attempt_id"] == control_attempt)
    assert effective["unparsed"] == []
    assert all(check["kind"] != "completeness" for check in effective["checks"])
    assert control_storage.latest_snapshots()[0]["unparsed"] == []
    assert control_storage.review_attempt(control_attempt)["preview"]["unparsed"]
    storage = Storage(tmp_path / "private")
    account = storage.create_account("sber", "brokerage", "Synthetic")
    iis_account = storage.create_account("sber", "iis", "Synthetic IIS")
    with pytest.raises(ValueError, match="contract kind"):
        storage.stage_file(iis_account, path)
    iis_report = tmp_path / "iis.html"
    iis_report.write_text(html.replace("Договор SYN123", "Договор на ведение индивидуального инвестиционного счета SYN123"))
    assert parse_sber(iis_report).account_kind == "iis"
    with pytest.raises(ValueError, match="contract kind"):
        storage.stage_file(account, iis_report)
    attempt1, preview1 = storage.stage_file(account, path)
    attempt2, preview2 = storage.stage_file(account, path)
    assert attempt2 != attempt1
    assert preview1 == preview2
    with storage._connect() as con:
        assert con.execute("SELECT COUNT(*) FROM source_files").fetchone()[0] == 1
        assert con.execute("SELECT COUNT(*) FROM ledger_events").fetchone()[0] == 0
        rows = con.execute("SELECT sheet,row_number,raw_json,classification FROM source_rows WHERE attempt_id=?", (attempt1,)).fetchall()
        assert any(row[0] == "table:2" and row[1] == 3 and row[3] == "cash_movement"
                   and json.loads(row[2])[2] == "Сделка от 15.01.2025" for row in rows)
        assert any(row[3] == "excluded" for row in rows)
    wrong_contract = tmp_path / "wrong-contract.html"
    wrong_contract.write_text(html.replace("SYN123", "OTHER456"))
    alias_storage = Storage(tmp_path / "alias-private")
    first_alias = alias_storage.create_account("sber", "brokerage", "First synthetic contract")
    second_alias = alias_storage.create_account("sber", "brokerage", "Second synthetic contract")
    first_alias_attempt, _ = alias_storage.stage_file(first_alias, path)
    assert alias_storage.confirm_attempt(first_alias_attempt) == {"inserted": 2, "observed_duplicates": 0}
    second_alias_attempt, _ = alias_storage.stage_file(second_alias, wrong_contract)
    alias_review = alias_storage.review_attempt(second_alias_attempt)
    assert any(item["status"] == "conflict" for item in alias_review["event_candidates"])
    assert any("another account" in item["reason"] and item["existing_options"]
               for item in alias_review["ambiguous_events"])
    assert len(alias_review["blocking_conflicts"]) == 1
    with pytest.raises(ValueError, match="already recorded under another account"):
        alias_storage.confirm_attempt(second_alias_attempt)
    assert alias_storage.review_attempt(second_alias_attempt)["state"] == "awaiting_review"
    assert alias_storage.list_events(second_alias) == []
    with alias_storage._connect() as con:
        assert con.execute("SELECT COUNT(*) FROM account_identities").fetchone()[0] == 1
        assert con.execute("SELECT COUNT(*) FROM ledger_events").fetchone()[0] == 2
    wrong_attempt, _ = storage.stage_file(account, wrong_contract)
    assert storage.confirm_attempt(attempt1) == {"inserted": 2, "observed_duplicates": 0}
    stored_trade = next(event for event in storage.list_events(account) if event["event_type"] == "trade")
    assert stored_trade["payload"]["accrued_interest"] == "0"
    with pytest.raises(ValueError, match="contract does not match"):
        storage.confirm_attempt(wrong_attempt)
    assert storage.review_attempt(wrong_attempt)["state"] == "awaiting_review"
    with pytest.raises(ValueError, match="contract does not match"):
        storage.stage_file(account, wrong_contract)
    with storage._connect() as con:
        con.execute("DROP TABLE account_identities")
        con.execute("PRAGMA user_version=5")
    storage = Storage(tmp_path / "private")
    with pytest.raises(ValueError, match="contract does not match"):
        storage.stage_file(account, wrong_contract)
    assert storage.list_instruments()[0]["isin"] == "RU000A000001"
    assert next(event for event in storage.list_events(account) if event["event_type"] == "trade")["instrument_id"] == storage.list_instruments()[0]["id"]
    assert storage.portfolio_estimate()["missing_prices"] == 1
    storage.add_price("RU000A000001", D("15"), date(2025, 1, 31), note="synthetic price")
    estimate = storage.portfolio_estimate()
    assert estimate["known_securities_subtotal"] == D("45")
    assert estimate["incomplete"] and estimate["missing_accounts"] == 1  # The empty IIS is not a proven zero.
    assert storage.confirm_attempt(attempt1) == {"inserted": 0, "observed_duplicates": 0}
    assert storage.confirm_attempt(attempt2) == {"inserted": 0, "observed_duplicates": 2}
    with storage._connect() as con:
        assert con.execute("SELECT COUNT(*) FROM ledger_events").fetchone()[0] == 2
        assert con.execute("SELECT COUNT(*) FROM event_observations").fetchone()[0] == 2
    backup = tmp_path / "backup.zip"
    create_backup(storage, backup)
    restored = restore_backup(backup, tmp_path / "restored")
    assert restored.list_events(account) == storage.list_events(account)
    with restored._connect() as con:
        source_path = con.execute("SELECT private_path FROM source_files ORDER BY id LIMIT 1").fetchone()[0]
    assert source_path.startswith(str(tmp_path / "restored"))

    overlap = tmp_path / "overlap.html"
    overlap.write_text(html.replace("Отчет брокера за период с", "Отчет брокера за период с\n"))
    overlap_attempt, _ = storage.stage_file(account, overlap)
    review = storage.review_attempt(overlap_attempt)
    assert len(review["ambiguous_events"]) == 1
    assert any(row["classification"] == "cash_movement" and row["raw"][2] == "Сделка от 15.01.2025"
               for row in review["diagnostic_rows"])
    try:
        storage.confirm_attempt(overlap_attempt)
    except ValueError as exc:
        assert "Ambiguous overlapping cash" in str(exc)
    else:
        raise AssertionError("overlap must require a decision")
    conflict = review["ambiguous_events"][0]
    result = storage.confirm_attempt(overlap_attempt, {conflict["key"]: {"action": "skip", "reason": "same synthetic cash row"}})
    assert result == {"inserted": 0, "observed_duplicates": 2}
    assert len(storage.list_events(account)) == 2
    renamed = tmp_path / "renamed.html"
    renamed.write_text(html.replace("Synthetic", "Renamed"))
    renamed_attempt, _ = storage.stage_file(account, renamed)
    renamed_review = storage.review_attempt(renamed_attempt)
    assert any(item["status"] == "duplicate" and item["key"].startswith("event:trade:")
               for item in renamed_review["event_candidates"])
    assert not any(item["reason"].startswith("same broker ID") for item in renamed_review["ambiguous_events"])
    renamed_decisions = {item["key"]: {"action": "skip", "reason": "same synthetic cash event"}
                         for item in renamed_review["ambiguous_events"]}
    assert storage.confirm_attempt(renamed_attempt, renamed_decisions) == {"inserted": 0, "observed_duplicates": 2}
    reject_attempt, _ = storage.stage_file(account, path)
    storage.reject_attempt(reject_attempt, "synthetic rejected reload")
    assert storage.review_attempt(reject_attempt)["state"] == "rejected"
    legacy_attempt, _ = storage.stage_file(account, path)
    with storage._connect() as con:
        con.execute("DROP TRIGGER protect_import_attempt_fields")
        con.execute("UPDATE import_attempts SET preview_json=json_remove(preview_json,'$.account_key') WHERE id=?", (legacy_attempt,))
    with pytest.raises(ValueError, match="Restage report"):
        storage.confirm_attempt(legacy_attempt)
    with storage._connect() as con:
        assert con.execute("SELECT COUNT(*) FROM ledger_events").fetchone()[0] == 2
        con.execute("DROP TRIGGER protect_account_identity")
        con.execute("DROP TABLE account_identities")
        con.execute("UPDATE import_attempts SET state='committed' WHERE id=?", (wrong_attempt,))
        con.execute("PRAGMA user_version=5")
    with pytest.raises(ValueError, match="multiple contracts"):
        Storage(tmp_path / "private")


def test_alfa_cash_period_summary_checks_both_boundaries(tmp_path):
    book = Workbook()
    book.remove(book.active)
    for name in ["Динамика позиций", "Завершенные сделки", "Незавершенные сделки", " Движение ДС", "Неторговые операции"]:
        sheet = book.create_sheet(name)
        sheet.cell(5, 8, "01.01.2025 - 31.01.2025")
    cash = book[" Движение ДС"]
    cash.cell(6, 9, "№ 12345678 от 01.01.2025")
    cash.cell(14, 12, "На начало периода")
    cash.cell(15, 15, "RUB")
    for col in (15, 17, 20):
        cash.cell(16, col, 50)
    cash.cell(18, 10, "Наименование операции")
    for col, value in {3: "15.01.2025", 10: "Комиссия", 15: -20}.items():
        cash.cell(20, col, value)
    cash.cell(21, 11, "Итого:")
    cash.cell(21, 15, -20)
    cash.cell(22, 12, "На конец периода")
    cash.cell(23, 15, "RUB")
    for col in (15, 17, 20):
        cash.cell(24, col, 30)
    pos = book["Динамика позиций"]
    pos.cell(19, 3, "Инструмент")
    pos.cell(19, 7, "Актив")
    for col, value in {3: "Валюта", 7: "RUB", 13: 50, 15: 50, 16: 30, 18: 30}.items():
        pos.cell(20, col, value)
    for col, value in {7: "Стоимость всех позиций, руб.", 15: 50, 18: 30}.items():
        pos.cell(22, col, value)
    path = tmp_path / "cash-summary.xlsx"
    book.save(path)
    report = parse_alfa(path)
    assert sum(row.reason == "Alfa period cash and portfolio summary; see source_total checks"
               for row in report.excluded) == 2
    assert sum(check.kind == "source_total" and check.status == "matched" for check in reconcile(report)) == 9
    cash.cell(24, 20, 31)
    book.save(path)
    assert any(check.kind == "source_total" and check.status == "difference"
               for check in reconcile(parse_alfa(path)))
    cash.cell(20, 3, "01.02.2025")
    book.save(path)
    outside = parse_alfa(path)
    assert not outside.cash_movements
    assert any(row.location.sheet == " Движение ДС" and row.location.row == 20
               and row.reason == "cash movement date outside report period" for row in outside.unparsed)
    cash.cell(20, 3, "not-a-date")
    book.save(path)
    invalid = parse_alfa(path)
    assert not invalid.cash_movements
    assert any(row.location.sheet == " Движение ДС" and row.location.row == 20
               and row.reason == "cash movement has invalid date" for row in invalid.unparsed)


def test_sber_cash_movement_outside_report_period_is_diagnostic_not_ledger_event(tmp_path):
    html = "Отчет брокера за период с 01.01.2025 по 31.01.2025 Договор SYN123 от 01.01.2025 " + "".join([
        _table([["Основной рынок", "Начало периода", "Конец периода"],
                ["Наименование", "ISIN ценной бумаги"] + ["-"] * 16]),
        _table([["Торговая площадка", "Валюта", "Конец периода"], ["1", "2", "3"],
                ["Торговый счет, Основной рынок", "RUB", "1", "0", "0", "0", "0", "0", "0"]]),
        _table([["Дата", "Торговая площадка", "Описание операции", "Валюта", "Сумма зачисления", "Сумма списания"],
                ["1", "2", "3", "4", "5", "6"],
                ["01.02.2025", "Основной рынок", "Пополнение", "RUB", "10", "0"]]),
    ])
    path = tmp_path / "outside.html"
    path.write_text(html)
    report = parse_sber(path)
    assert not report.cash_movements
    assert any(row.location.sheet == "table:2" and row.location.row == 3
               and row.reason == "cash: cash movement date outside report period" for row in report.unparsed)
    store = Storage(tmp_path / "private")
    account = store.create_account("sber", "brokerage", "Synthetic")
    attempt_id, _ = store.stage_file(account, path)
    preview = store.review_attempt(attempt_id)
    assert not preview["preview"]["cash_movements"]
    assert any(row["row"] == 3 and row["sheet"] == "table:2" for row in preview["diagnostic_rows"])


def test_alfa_xlsx_report_reconciles(tmp_path):
    book = Workbook()
    book.remove(book.active)
    for name in ["Динамика позиций", "Завершенные сделки", "Незавершенные сделки", " Движение ДС", "Неторговые операции"]:
        sheet = book.create_sheet(name)
        sheet.cell(5, 8, "01.01.2025 - 31.01.2025")
    trade = book["Завершенные сделки"]
    for col, value in {5: "№ сделки", 13: "ISIN/рег.код", 22: "Вал.", 23: "Комиссия банка"}.items():
        trade.cell(15, col, value)
    for col, value in {5: "deal-1", 8: "15.01.2025\n 12:00:00", 11: "15.01.2025", 13: "RU000A000001", 15: "Synthetic", 17: 2, 19: 10, 20: 20, 22: "RUB", 23: 0}.items():
        trade.cell(16, col, value)
    trade.cell(17, 12, "⁶ Места заключения сделок: синтетическое пояснение")
    cash = book[" Движение ДС"]
    cash.cell(6, 9, "№ 12345678 от 01.01.2025")
    cash.cell(18, 10, "Наименование операции")
    for col, value in {3: "15.01.2025", 10: "Расчеты по сделке deal-1", 15: -20}.items():
        cash.cell(20, col, value)
    cash.cell(21, 11, "Итого:")
    cash.cell(21, 15, -20)
    pos = book["Динамика позиций"]
    pos.cell(19, 3, "Инструмент")
    pos.cell(19, 7, "Актив")
    for col, value in {3: "Валюта", 7: "RUB", 13: 50, 16: 30}.items():
        pos.cell(20, col, value)
    for col, value in {3: "Акции", 7: "Synthetic", 13: 1, 16: 3}.items():
        pos.cell(22, col, value)
    pos.cell(23, 7, "RU000A000001")
    for col, value in {7: "Continuation security", 13: 5, 16: 5}.items():
        pos.cell(24, col, value)
    pos.cell(25, 7, "RU000A000002")
    pos.cell(26, 7, "Стоимость всех позиций, руб.")
    pos.cell(26, 15, 0)
    pos.cell(26, 18, 0)
    for row, label, amount in ((28, "НДФЛ исчислено, руб.", 20),
                               (29, "Удержано, руб.", 5), (30, "К доплате, руб.", 15)):
        pos.cell(row, 15, label)
        pos.cell(row, 18, amount)
    pos.cell(27, 15, "Расчет обязательств по НДФЛ")
    nontrade = book["Неторговые операции"]
    nontrade.cell(15, 10, "Наименование операции")
    nontrade.cell(20, 1, "Периоды использования единого финансового рынка")
    path = tmp_path / "report.xlsx"
    book.save(path)
    report = parse_alfa(path)
    assert report.variant == "alfa-xls-dynamic-positions-v12"
    assert report.trades[0].execution_time == "12:00:00"
    trade.cell(16, 11, "01.02.2025")
    book.save(path)
    outside = parse_alfa(path)
    assert not outside.trades
    assert any(row.location.sheet == "Завершенные сделки" and row.location.row == 16
               and "completed trade settlement outside report period" in row.reason for row in outside.unparsed)
    trade.cell(16, 11, "14.01.2025")
    book.save(path)
    invalid_order = parse_alfa(path)
    assert not invalid_order.trades
    assert any("trade settlement precedes execution" in row.reason for row in invalid_order.unparsed)
    trade.cell(16, 11, "15.01.2025")
    trade.cell(16, 8, "15.01.2025")
    book.save(path)
    assert any("trade execution time is absent or ambiguous" in row.reason
               for row in parse_alfa(path).unparsed)
    trade.cell(16, 8, "15.01.2025\n 12:00:00")
    book.save(path)
    assert any(row.reason == "Alfa market access period metadata; no security operation"
               for row in report.excluded)
    nontrade.cell(27, 1, "Synthetic unknown note")
    book.save(path)
    assert any(row.reason == "unrecognized nonempty Alfa source row"
               for row in parse_alfa(path).unparsed)
    nontrade.cell(27, 1).value = None
    nontrade.cell(27, 3, "15.01.2025")
    nontrade.cell(27, 10, "Synthetic security transfer")
    book.save(path)
    assert any(row.reason == "non-trading security operation requires classification"
               for row in parse_alfa(path).unparsed)
    nontrade.cell(27, 3).value = None
    nontrade.cell(27, 10).value = None
    book.save(path)
    report = parse_alfa(path)
    assert any(row.reason == "Alfa empty trade sheet heading; no ledger event" for row in report.excluded)
    assert any(row.reason == "Alfa trade table footer note; no ledger event" for row in report.excluded)
    assert len(report.source_totals) == 4
    pos.cell(30, 18, 16)
    book.save(path)
    assert any(check.kind == "source_total" and check.key.endswith("tax_due") and check.status == "difference"
               for check in reconcile(parse_alfa(path)))
    pos.cell(30, 18, 15)
    book.save(path)
    pos.cell(26, 18, 1)
    book.save(path)
    assert any(check.kind == "source_total" and check.status == "difference"
               for check in reconcile(parse_alfa(path)))
    pos.cell(26, 18, 0)
    book.save(path)
    report = parse_alfa(path)
    assert sum(row.reason == "ISIN continuation used by preceding position; no ledger event"
               for row in report.excluded) == 2
    assert any(row.reason == "Alfa report heading or column labels; no ledger event" for row in report.excluded)
    assert report.account_key is not None
    assert report.account_number == "12345678"
    assert report.trades[0].quantity == D("2")
    assert {item.isin for item in report.positions} == {"RU000A000001", "RU000A000002"}
    assert all(check.status == "matched" for check in reconcile(report))
    assert check_alfa(path, report)["cash_semantic_source_row"] == 1
    report.cash_movements[0] = replace(report.cash_movements[0], kind="deposit")
    with pytest.raises(ValueError, match="cash category"):
        check_alfa(path, report)
    report = parse_alfa(path)
    storage = Storage(tmp_path / "private")
    account = storage.create_account("alfa", "brokerage", "Synthetic")
    first_attempt, _ = storage.stage_file(account, path)
    assert storage.review_attempt(first_attempt)["structural_rows"]
    with storage._connect() as con:
        raw = con.execute("SELECT raw_json FROM source_rows WHERE attempt_id=? AND sheet=? AND row_number=? AND classification='cash_movement'",
                          (first_attempt, " Движение ДС", 20)).fetchone()
    assert json.loads(raw[0])[9] == "Расчеты по сделке deal-1"
    assert storage.confirm_attempt(first_attempt) == {"inserted": 2, "observed_duplicates": 0}
    trade.cell(1, 1, "second file with the same operations")
    overlapping = tmp_path / "overlap.xlsx"
    book.save(overlapping)
    second_attempt, _ = storage.stage_file(account, overlapping)
    assert storage.review_attempt(second_attempt)["ambiguous_events"] == []
    assert storage.confirm_attempt(second_attempt) == {"inserted": 0, "observed_duplicates": 2}
    cash.cell(20, 15, -21)
    cash.cell(21, 15, -21)
    pos.cell(20, 16, 29)
    changed = tmp_path / "changed.xlsx"
    book.save(changed)
    changed_attempt, _ = storage.stage_file(account, changed)
    assert any(item["reason"] == "same broker ID has different payload"
               for item in storage.review_attempt(changed_attempt)["ambiguous_events"])
    with pytest.raises(ValueError, match="Conflicting broker event ID"):
        storage.confirm_attempt(changed_attempt)
    cash.cell(20, 15, -20)
    cash.cell(21, 15, -20)
    pos.cell(20, 16, 30)
    for col, value in {3: "Валюта", 7: "USD", 13: 10, 16: 10}.items():
        pos.cell(21, col, value)
    book.save(path)
    assert any(row.reason == "non-RUB cash balance" for row in parse_alfa(path).unparsed)


def test_alfa_income_transfer_requires_matching_report_subtotal(tmp_path):
    book = Workbook()
    book.remove(book.active)
    for name in ["Динамика позиций", "Завершенные сделки", "Незавершенные сделки", " Движение ДС", "Неторговые операции"]:
        sheet = book.create_sheet(name)
        sheet.cell(5, 8, "01.01.2025 - 31.01.2025")
    cash = book[" Движение ДС"]
    cash.cell(6, 9, "№ 12345678 от 01.01.2025")
    cash.cell(18, 10, "Наименование операции")
    for row, comment, amount in ((20, "Дивиденды за 2024 г. по акциям Synthetic (Удержан налог 13 руб.)", 100),
                                 (21, "погашение купона Synthetic", 50)):
        cash.cell(row, 3, "15.01.2025")
        cash.cell(row, 10, "Перевод")
        cash.cell(row, 11, comment)
        cash.cell(row, 15, amount)
    cash.cell(22, 11, "Итого:")
    cash.cell(22, 15, 150)
    cash.cell(27, 12, "Дивиденды")
    cash.cell(27, 15, 100)
    cash.cell(28, 12, "Купоны")
    cash.cell(28, 15, 50)
    positions = book["Динамика позиций"]
    positions.cell(19, 3, "Инструмент")
    positions.cell(19, 7, "Актив")
    for col, value in {3: "Валюта", 7: "RUB", 13: 0, 16: 150}.items():
        positions.cell(20, col, value)
    path = tmp_path / "income.xlsx"
    book.save(path)
    report = parse_alfa(path)
    assert [item.kind for item in report.cash_movements] == ["dividend", "tax_withheld", "coupon"]
    assert [item.amount for item in report.cash_movements] == [D("113"), D("-13"), D("50")]
    assert report.cash_movements[0].withheld_tax == D("13")
    assert report.cash_movements[1].component == "dividend_tax"
    assert [item.kind for item in report.cash_flow_subtotals] == ["dividend", "coupon"]
    assert [check.status for check in reconcile(report) if check.kind == "cash_flow"] == ["matched", "matched"]
    assert not report.unparsed
    assert next(check for check in reconcile(report) if check.kind == "cash").status == "matched"
    store = Storage(tmp_path / "income-private")
    account = store.create_account("alfa", "brokerage", "Synthetic income")
    attempt, _ = store.stage_file(account, path)
    with store._connect() as con:
        assert con.execute("SELECT COUNT(*) FROM source_rows WHERE attempt_id=? AND classification='cash_flow_subtotal'",
                           (attempt,)).fetchone()[0] == 2
        assert con.execute("SELECT COUNT(*) FROM source_rows WHERE attempt_id=? AND classification='cash_movement'",
                           (attempt,)).fetchone()[0] == 2
    review_decisions = {key: {"action": "acknowledge", "reason": "synthetic fixture reviewed"}
                        for key in store.review_attempt(attempt)["required_decisions"]}
    assert store.confirm_attempt(attempt, review_decisions) == {"inserted": 3, "observed_duplicates": 0}
    replay, _ = store.stage_file(account, path)
    replay_decisions = {key: {"action": "acknowledge", "reason": "synthetic fixture reviewed"}
                        for key in store.review_attempt(replay)["required_decisions"]}
    assert store.confirm_attempt(replay, replay_decisions) == {"inserted": 0, "observed_duplicates": 3}
    assert sum((D(event["payload"]["amount"]) for event in store.list_events(account)
                if event["event_type"] == "cash"), D(0)) == D("150")

    legacy = Storage(tmp_path / "legacy-income-private")
    legacy_account = legacy.create_account("alfa", "brokerage", "Synthetic legacy income")
    legacy_attempt, _ = legacy.stage_file(legacy_account, path)
    with legacy._connect() as con:
        source_id = con.execute("SELECT source_id FROM import_attempts WHERE id=?", (legacy_attempt,)).fetchone()[0]
        con.execute("DROP TRIGGER validate_ledger_source_row")
        con.execute("DROP TRIGGER validate_observation_source_row")
        old_event = con.execute(
            "INSERT INTO ledger_events(account_id,source_id,sheet,row_number,event_type,broker_event_id,fingerprint,payload_json,created_at) "
            "VALUES(?,?,?,?,?,?,?,?,?)",
            (legacy_account, source_id, " Движение ДС", 20, "cash", None, "old-net-dividend",
             json.dumps({"kind": "dividend", "date": "2025-01-15", "amount": "100",
                         "withheld_tax": "13", "currency": "RUB"}), "2025-01-15"),
        ).lastrowid
        con.execute("INSERT INTO event_observations(event_id,source_id,sheet,row_number,event_type) VALUES(?,?,?,?,?)",
                    (old_event, source_id, " Движение ДС", 20, "cash"))
    legacy_decisions = {key: {"action": "acknowledge", "reason": "synthetic legacy fixture reviewed"}
                        for key in legacy.review_attempt(legacy_attempt)["required_decisions"]}
    with pytest.raises(ValueError, match="Previously imported source row"):
        legacy.confirm_attempt(legacy_attempt, legacy_decisions)
    assert len(legacy.list_events(legacy_account)) == 1

    cash.cell(27, 15, 99)
    book.save(path)
    mismatched = parse_alfa(path)
    assert [item.kind for item in mismatched.cash_movements] == ["transfer_unclassified", "coupon"]
    assert any("matching report subtotal" in item.reason for item in mismatched.unparsed)
    assert any(check.kind == "cash_flow" and check.key == "dividend" and check.status == "difference"
               for check in reconcile(mismatched))
    cash.cell(27, 15, 100)
    cash.cell(20, 11, "Дивиденды за 2024 г. по акциям Synthetic")
    book.save(path)
    assert any("withholding amount is absent" in item.reason for item in parse_alfa(path).unparsed)


def test_alfa_bank_tariff_subtotal_checks_both_fee_types(tmp_path):
    book = Workbook()
    book.remove(book.active)
    for name in ["Динамика позиций", "Завершенные сделки", "Незавершенные сделки", " Движение ДС", "Неторговые операции"]:
        sheet = book.create_sheet(name)
        sheet.cell(5, 8, "01.01.2025 - 31.01.2025")
    cash = book[" Движение ДС"]
    cash.cell(6, 9, "№ 12345678 от 01.01.2025")
    cash.cell(18, 10, "Наименование операции")
    for row, label, amount in ((20, "Комиссия", -3), (21, "Комиссия по сделке 123", -2)):
        cash.cell(row, 3, "15.01.2025")
        cash.cell(row, 10, label)
        cash.cell(row, 15, amount)
    cash.cell(22, 11, "Итого:")
    cash.cell(30, 12, "Списано по тарифам Банка")
    cash.cell(30, 15, -5)
    positions = book["Динамика позиций"]
    positions.cell(19, 3, "Инструмент")
    positions.cell(19, 7, "Актив")
    for col, value in {3: "Валюта", 7: "RUB", 13: 10, 16: 5}.items():
        positions.cell(20, col, value)
    path = tmp_path / "fees.xlsx"
    book.save(path)
    report = parse_alfa(path)
    assert [(row.kind, row.amount) for row in report.cash_movements] == [("fee", D("-3")), ("trade_fee", D("-2"))]
    assert [(row.kind, row.amount) for row in report.cash_flow_subtotals] == [("bank_fee", D("-5"))]
    assert next(check for check in reconcile(report) if check.kind == "cash_flow").status == "matched"
    cash.cell(30, 15, -4)
    book.save(path)
    assert next(check for check in reconcile(parse_alfa(path)) if check.kind == "cash_flow").status == "difference"


def test_alfa_transfer_comments_distinguish_external_cash_and_unknown_counterparty(tmp_path):
    book = Workbook()
    book.remove(book.active)
    for name in ["Динамика позиций", "Завершенные сделки", "Незавершенные сделки", " Движение ДС", "Неторговые операции"]:
        sheet = book.create_sheet(name)
        sheet.cell(5, 8, "01.01.2025 - 31.01.2025")
    cash = book[" Движение ДС"]
    cash.cell(6, 9, "№ 12345678 от 01.01.2025")
    cash.cell(18, 10, "Наименование операции")
    for row, comment, amount in ((20, 'из АО "Альфа-Банк"', 100),
                                 (21, "Возврат средств по дог. № SYN", -30),
                                 (22, "с субсчета 999-000", 20),
                                 (23, 'из АО "Альфа-Банк"', -1)):
        cash.cell(row, 3, "15.01.2025")
        cash.cell(row, 10, "Перевод")
        cash.cell(row, 11, comment)
        cash.cell(row, 15, amount)
    cash.cell(24, 11, "Итого:")
    cash.cell(24, 15, 89)
    cash.cell(28, 12, "Внесено по распоряжению Клиента")
    cash.cell(28, 15, 120)
    cash.cell(29, 12, "Списано по распоряжению Клиента")
    cash.cell(29, 15, -30)
    positions = book["Динамика позиций"]
    positions.cell(19, 3, "Инструмент")
    positions.cell(19, 7, "Актив")
    for col, value in {3: "Валюта", 7: "RUB", 13: 0, 16: 89}.items():
        positions.cell(20, col, value)
    path = tmp_path / "transfers.xlsx"
    book.save(path)
    report = parse_alfa(path)
    assert [item.kind for item in report.cash_movements] == [
        "deposit", "withdrawal", "deposit", "transfer_unclassified"]
    assert {item.location.row for item in report.unparsed} == {22, 23}
    assert next(check for check in reconcile(report) if check.kind == "cash").status == "matched"
    assert [check.status for check in reconcile(report) if check.kind == "cash_flow"] == ["matched", "matched"]
    store = Storage(tmp_path / "subtotal-private")
    account = store.create_account("alfa", "brokerage", "Synthetic")
    attempt, preview = store.stage_file(account, path)
    assert len(preview["cash_flow_subtotals"]) == 2
    with store._connect() as con:
        assert con.execute("SELECT COUNT(*) FROM source_rows WHERE attempt_id=? AND classification='cash_flow_subtotal'",
                           (attempt,)).fetchone()[0] == 2
    cash.cell(28, 15, 119)
    book.save(path)
    assert any(check.kind == "cash_flow" and check.status == "difference" for check in reconcile(parse_alfa(path)))


def test_confirmation_blocks_unknown_rows_and_preserves_journal(tmp_path):
    # Same synthetic report shape as above, with one unknown cash operation.
    base = "Отчет брокера за период с 01.01.2025 по 31.01.2025 Договор SYN123 от 01.01.2025 "
    html = base + "".join([
        _table([["Основной рынок", "Начало периода", "Конец периода"],
                ["Наименование", "ISIN ценной бумаги"] + ["-"] * 16,
                ["Synthetic", "RU000A000001", "RUB", "1"] + ["0"] * 4 + ["1"] + ["0"] * 9]),
        _table([["Торговая площадка", "Валюта", "Конец периода"], ["1", "2", "3"],
                ["Торговый счет, Основной рынок", "RUB", "1", "50", "0", "50", "0", "0", "50"]]),
        _table([["Дата", "Торговая площадка", "Описание операции", "Валюта", "Сумма зачисления", "Сумма списания"],
                ["1", "2", "3", "4", "5", "6"], ["15.01.2025", "Основной рынок", "Неизвестная операция", "RUB", "1", ""]]),
        _table([["Наименование", "Код", "ISIN ценной бумаги"], ["1", "2", "3"], ["Synthetic", "SYN", "RU000A000001"]]),
    ])
    path = tmp_path / "unknown.html"
    path.write_text(html)
    storage = Storage(tmp_path / "private")
    account = storage.create_account("sber", "brokerage", "Synthetic")
    attempt, _ = storage.stage_file(account, path)
    review = storage.review_attempt(attempt)
    assert review["required_decisions"]
    assert any(row["classification"] == "unparsed" and row["raw"][2] == "Неизвестная операция"
               for row in review["diagnostic_rows"])
    with pytest.raises(ValueError, match="Explicit review"):
        storage.confirm_attempt(attempt)
    assert storage.list_events(account) == []


def test_trade_payload_name_is_not_economic_identity():
    old = '{"broker_id":"deal-1","isin":"RU000A000001","name":"Old","quantity":"1","amount":"10"}'
    renamed = '{"broker_id":"deal-1","isin":"RU000A000001","name":"New","quantity":"1","amount":"10"}'
    changed_amount = '{"broker_id":"deal-1","isin":"RU000A000001","name":"New","quantity":"1","amount":"11"}'
    assert Storage._same_event_payload("trade", old, renamed)
    assert not Storage._same_event_payload("trade", old, changed_amount)
    assert Storage._same_event_payload("trade", old.replace('"1"', '"1.0"').replace('"10"', '"10.00"'), renamed)
    assert not Storage._same_event_payload("trade", old, renamed.replace('"amount":"10"',
                                                                           '"amount":"10","accrued_interest":"0"'))
    assert Storage._same_event_payload("cash", '{"amount":"-20.0","kind":"trade_cash"}',
                                       '{"amount":"-20","kind":"trade_cash"}')


def test_repeated_cash_reference_within_one_report_stays_unkeyed():
    movement = {"kind": "trade_fee", "broker_trade_id": "deal-1", "date": "2025-01-01",
                "amount": "-1", "currency": "RUB"}
    preview = {"trades": [], "cash_movements": [
        {**movement, "location": {"sheet": "cash", "row": 1}},
        {**movement, "location": {"sheet": "cash", "row": 2}},
    ]}
    assert all(item["broker_id"] is None for item in Storage._event_items(preview))


def test_sber_malformed_dated_movement_is_visible(tmp_path):
    html = "Отчет брокера за период с 01.01.2025 по 31.01.2025 Договор SYN123 от 01.01.2025 " + "".join([
        _table([["Основной рынок", "Начало периода", "Конец периода"],
                ["Наименование", "ISIN ценной бумаги"] + ["-"] * 16,
                ["Synthetic", "RU000A000001", "RUB", "0"] + ["0"] * 14]),
        _table([["Торговая площадка", "Валюта", "Конец периода"], ["1", "2", "3"],
                ["Торговый счет, Основной рынок", "RUB", "1", "0", "0", "0", "0", "0", "0"]]),
        _table([["Дата", "Торговая площадка", "Описание операции", "Валюта", "Сумма зачисления", "Сумма списания"],
                ["1", "2", "3", "4", "5", "6"], ["15.01.2025", "Основной рынок", "Пополнение", "RUB", "10"]]),
    ])
    path = tmp_path / "malformed.html"
    path.write_text(html)
    report = parse_sber(path)
    assert not report.cash_movements
    assert any("unexpected column count" in row.reason for row in report.unparsed)


def test_alfa_trade_without_broker_id_is_visible(tmp_path):
    book = Workbook()
    book.remove(book.active)
    for name in ["Динамика позиций", "Завершенные сделки", "Незавершенные сделки", " Движение ДС", "Неторговые операции"]:
        sheet = book.create_sheet(name)
        sheet.cell(5, 8, "01.01.2025 - 31.01.2025")
    trade = book["Завершенные сделки"]
    trade.cell(15, 5, "№ сделки")
    trade.cell(15, 13, "ISIN/рег.код")
    trade.cell(16, 13, "RU000A000001")
    trade.cell(16, 17, 2)
    cash = book[" Движение ДС"]
    cash.cell(6, 9, "№ 12345678 от 01.01.2025")
    cash.cell(18, 10, "Наименование операции")
    cash.cell(20, 3, "15.01.2025")
    cash.cell(20, 7, "10:00")
    cash.cell(20, 15, 10)
    positions = book["Динамика позиций"]
    positions.cell(19, 3, "Инструмент")
    positions.cell(19, 7, "Актив")
    path = tmp_path / "malformed.xlsx"
    book.save(path)
    report = parse_alfa(path)
    assert not report.trades
    assert any("without broker ID" in row.reason for row in report.unparsed)
    assert any("cash amount without operation type" in row.reason for row in report.unparsed)


def test_report_without_unique_contract_is_rejected(tmp_path):
    html = tmp_path / "no-contract.html"
    html.write_text("Отчет брокера за период с 01.01.2025 по 31.01.2025")
    with pytest.raises(UnsupportedReport, match="contract identity"):
        parse_sber(html)

    book = Workbook()
    book.remove(book.active)
    for name in ["Динамика позиций", "Завершенные сделки", "Незавершенные сделки", " Движение ДС", "Неторговые операции"]:
        sheet = book.create_sheet(name)
        sheet.cell(5, 8, "01.01.2025 - 31.01.2025")
    xlsx = tmp_path / "no-agreement.xlsx"
    book.save(xlsx)
    with pytest.raises(UnsupportedReport, match="agreement identity"):
        parse_alfa(xlsx)
